#!/usr/bin/env python3
"""Standalone market-risk RWA service. Run from this directory or as a script path."""

from __future__ import annotations

import io
import os
import re
import shutil
import sys
import tempfile
import uuid
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))


def _load_local_env() -> None:
    for env_path in (APP_DIR.parent / ".env", APP_DIR / ".env"):
        if not env_path.is_file():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            key, val = key.strip(), val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val


_load_local_env()

from flask import Flask, jsonify, request, send_file

from agent.subagents.rwa.market.agent import field_catalog, inspect_workbook, list_type_columns, run_with_mappings, write_rwa_workbook
from agent.subagents.rwa.market.compare import compare_module, module_catalog, write_comparison
from agent.subagents.weight.calc.exposure_class import prototype_flow
from agent.subagents.weight.paper.paper_calc import load_or_compute, price_open_facts, write_annotated_source
from agent.subagents.weight.paper.working_paper import load_working_paper
from knowledge.local.library import calculation_rules, ensure_library, load_library, reset_rule, save_rule
from agent.subagents.rwa.market.rule_calc import run_authored
from knowledge.local.rule_graph import load_graph, save_graph
from agent.subagents.weight.paper.exp_flow_gen import analyze_flow_gaps_local, generate_product_flow
from agent.subagents.rwa.market.mapper import suggest_field_map, suggest_mappings
from agent.platform.graph import reply as agent_reply
from agent.tree.edit import propose as propose_tree
from knowledge.weknora.rule_chat import discuss_rules, lookup_basis
from knowledge.local.rule_learn import learn_from_excel
from knowledge.weknora.weknora import ask as weknora_ask
from knowledge.weknora.weknora import list_bases as weknora_bases
from agent.subagents.rwa.market.weight_tables import ingest_weight_workbook, weight_catalog

ensure_library()

MAX_UPLOAD_MB = 120

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

# upload_id -> (workbook path, work dir, original filename)
UPLOADS: dict[str, tuple[Path, Path, str]] = {}
# upload_id -> exposure grid from the latest detail calculation
EXPOSURES: dict[str, dict] = {}


def _ingest_weights(path: Path, filename: str) -> dict:
    try:
        return ingest_weight_workbook(path, filename)
    except Exception:
        return {"added": False, "duplicate": False, "filename": filename}


def _result_suffix(filename: str) -> str:
    return ".xlsm" if filename.lower().endswith(".xlsm") else ".xlsx"


def _result_name(filename: str) -> str:
    stem = Path(filename).stem or "结果"
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", stem).strip()[:80] or "结果"
    return f"{stem}_计算RWA{_result_suffix(filename)}"


def _result_path(work_dir: Path, filename: str) -> Path:
    return work_dir / ("calculated" + _result_suffix(filename))


@app.get("/health")
def health():
    return jsonify({"status": "ok", "service": "rwa"})


def _weixin_env_path() -> Path:
    return APP_DIR / ".env"


def _upsert_env(updates: dict[str, str], path: Path | None = None) -> None:
    path = path or _weixin_env_path()
    lines = path.read_text(encoding="utf-8").splitlines() if path.is_file() else []
    seen: set[str] = set()
    written: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            written.append(line)
            continue
        key = stripped.split("=", 1)[0].strip()
        if key in updates:
            written.append(f"{key}={updates[key]}")
            seen.add(key)
        else:
            written.append(line)
    for key, value in updates.items():
        if key not in seen:
            written.append(f"{key}={value}")
    path.write_text("\n".join(written) + "\n", encoding="utf-8")
    for key, value in updates.items():
        os.environ[key] = value


def _credential(value: object, limit: int) -> str:
    text = str(value or "").strip()
    if any(char in text for char in "\r\n"):
        raise ValueError("不能包含换行")
    if len(text) > limit:
        raise ValueError("内容过长")
    return text


@app.get("/api/weixin")
def weixin_get():
    appid = os.environ.get("WEIXIN_APPID", "").strip()
    secret = os.environ.get("WEIXIN_APPSECRET", "").strip()
    return jsonify({"appid": appid, "secret_set": bool(secret)})


@app.put("/api/weixin")
def weixin_save():
    data = request.get_json(silent=True) or {}
    try:
        appid = _credential(data.get("appid"), 64)
        secret = _credential(data.get("secret"), 128)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    if not appid:
        return jsonify({"error": "请填写 AppID"}), 400
    updates = {"WEIXIN_APPID": appid}
    if secret:
        updates["WEIXIN_APPSECRET"] = secret
    elif not os.environ.get("WEIXIN_APPSECRET", "").strip():
        return jsonify({"error": "请填写 AppSecret"}), 400
    _upsert_env(updates)
    return jsonify({"appid": appid, "secret_set": bool(os.environ.get("WEIXIN_APPSECRET", "").strip())})


@app.get("/api/fields")
def fields():
    return jsonify({"packs": field_catalog()})


@app.post("/api/inspect")
def inspect():
    upload = request.files.get("file")
    if upload is None or upload.filename == "":
        return jsonify({"error": "请先选择要上传的 Excel 文件"}), 400
    name = upload.filename or ""
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "仅支持 .xlsx 或 .xlsm 格式"}), 400
    work_dir = Path(tempfile.mkdtemp(prefix="rwa_map_"))
    input_path = work_dir / "input.xlsx"
    upload.save(input_path)
    upload_id = uuid.uuid4().hex
    UPLOADS[upload_id] = (input_path, work_dir, name)
    try:
        preview = inspect_workbook(input_path)
        preview["weights"] = _ingest_weights(input_path, name)
    except Exception as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        UPLOADS.pop(upload_id, None)
        return jsonify({"error": f"读取失败: {e}"}), 500
    preview["upload_id"] = upload_id
    preview["filename"] = name
    return jsonify(preview)


@app.post("/api/types")
def types():
    data = request.get_json(silent=True) or {}
    upload_id = str(data.get("upload_id") or "")
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "请重新上传文件"}), 400
    sheet = str(data.get("sheet") or "")
    try:
        header_row = int(data.get("header_row"))
    except (TypeError, ValueError):
        return jsonify({"error": "表头行无效"}), 400
    try:
        columns = list_type_columns(stored[0], sheet, header_row)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"读取失败: {e}"}), 500
    return jsonify({"type_columns": columns})


@app.post("/api/suggest")
def suggest():
    data = request.get_json(silent=True) or {}
    upload_id = str(data.get("upload_id") or "")
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "请重新上传文件"}), 400
    products = data.get("products")
    if not isinstance(products, list) or not products:
        return jsonify({"error": "请先选择计算方式"}), 400
    sheet = str(data.get("sheet") or "")
    try:
        header_row = int(data.get("header_row"))
    except (TypeError, ValueError):
        return jsonify({"error": "表头行无效"}), 400
    try:
        mappings = suggest_mappings(stored[0], sheet, header_row, [str(item) for item in products])
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    return jsonify({"mappings": mappings})


@app.post("/api/exp/generate-flow")
def exp_generate_flow():
    """通用流程模板 + 产品字段 → 生成产品分类流程，并标出字段缺口。"""
    data = request.get_json(silent=True) or {}
    label = str(data.get("product_label") or data.get("label") or "").strip()
    fields = data.get("product_fields") or data.get("fields") or []
    base_flow = data.get("base_flow") or data.get("template")
    if not label:
        return jsonify({"error": "缺少产品类型名称"}), 400
    if not isinstance(fields, list):
        return jsonify({"error": "product_fields 须为数组"}), 400
    if not isinstance(base_flow, dict) or not base_flow.get("categories"):
        return jsonify({"error": "缺少通用流程模板 base_flow"}), 400
    try:
        result = generate_product_flow(label, fields, base_flow)
    except Exception as e:
        # 模型不可用时回退本地缺口分析，保证原型可演示
        try:
            result = analyze_flow_gaps_local(label, fields, base_flow)
            result["fallback_error"] = str(e)
        except Exception as e2:
            return jsonify({"error": f"生成失败: {e2}"}), 500
    return jsonify(result)


@app.get("/api/working-paper")
def working_paper():
    """把底稿明细交给模型对上每一个分类字段，再算出风险加权资产。"""
    try:
        data = load_or_compute()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        return jsonify({"error": "底稿计算失败: " + str(exc)}), 500
    data["download_url"] = "/api/working-paper.xlsx"
    data["download_name"] = "信用风险RWA计算结果.xlsx"
    return jsonify(data)


@app.get("/api/working-paper.xlsx")
def working_paper_file():
    """计算完成后的结果表：汇总、分类、分表和字段对应。"""
    try:
        data = load_or_compute()
        buf = io.BytesIO()
        write_annotated_source(data, buf)
        buf.seek(0)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        return jsonify({"error": "结果导出失败: " + str(exc)}), 500
    return send_file(
        buf,
        as_attachment=True,
        download_name="信用风险RWA计算结果.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


CLASS_DOC_PATH = (
    APP_DIR / "agent" / "subagents" / "weight" / "knowledge" / "权重法风险暴露分类.json"
)


@app.get("/api/exposure-class.json")
def exposure_class_json():
    """分类流程图直接读取的规则原文。"""
    if not CLASS_DOC_PATH.is_file():
        return jsonify({"error": "找不到权重法风险暴露分类.json"}), 404
    response = send_file(CLASS_DOC_PATH, mimetype="application/json")
    response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/api/exposure-class")
def exposure_class():
    """原型分类画布读取的判定树。顺序和权重来自知识库 JSON。"""
    try:
        flow = prototype_flow()
    except Exception as e:
        return jsonify({"error": f"读取风险暴露分类失败: {e}"}), 500
    return jsonify({"flow": flow})


@app.get("/api/weights")
def weights():
    return jsonify(weight_catalog())


@app.get("/api/rules")
def rules_library():
    try:
        library = load_library()
    except Exception as e:
        return jsonify({"error": f"读取规则失败: {e}"}), 500
    return jsonify({"rules": library})


@app.get("/api/rules/graph")
def rules_graph_get():
    return jsonify({"graph": load_graph()})


@app.put("/api/rules/graph")
def rules_graph_save():
    data = request.get_json(silent=True) or {}
    try:
        graph = save_graph(data.get("graph"))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"写入图谱失败: {e}"}), 500
    return jsonify({"graph": graph})


@app.put("/api/rules")
def rules_save():
    data = request.get_json(silent=True) or {}
    pack = str(data.get("id") or "")
    try:
        if data.get("reset"):
            rule = reset_rule(pack)
        else:
            rule = save_rule(pack, data.get("rule"))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"写入规则失败: {e}"}), 500
    return jsonify({"rule": rule})


@app.post("/api/rules/learn")
def rules_learn():
    upload = request.files.get("file")
    if upload is None or upload.filename == "":
        return jsonify({"error": "请先选择规则 Excel"}), 400
    name = upload.filename or ""
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "规则 Excel 仅支持 .xlsx 或 .xlsm"}), 400
    upload.seek(0, os.SEEK_END)
    size = upload.tell()
    upload.seek(0)
    if size > 15 * 1024 * 1024:
        return jsonify({"error": "规则 Excel 请小于 15MB"}), 400
    work_dir = Path(tempfile.mkdtemp(prefix="rwa_rule_"))
    input_path = work_dir / "rules.xlsx"
    try:
        upload.save(input_path)
        weights = _ingest_weights(input_path, name)
        result = learn_from_excel(input_path, name)
        result["weights"] = weights
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
    return jsonify(result)


@app.get("/api/weknora/bases")
def weknora_list():
    try:
        return jsonify({"bases": weknora_bases()})
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 502


@app.post("/api/weknora/ask")
def weknora_query():
    data = request.get_json(silent=True) or {}
    try:
        result = weknora_ask(str(data.get("query") or ""), str(data.get("knowledge_base_id") or ""))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 502
    return jsonify(result)


@app.post("/api/rules/basis")
def rules_basis():
    data = request.get_json(silent=True) or {}
    try:
        basis = lookup_basis(data.get("title"))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    return jsonify({"basis": basis})


@app.post("/api/exposure-tree/propose")
def exposure_tree_propose():
    data = request.get_json(silent=True) or {}
    try:
        result = propose_tree(
            data.get("messages") or [],
            data.get("document") or {},
            data.get("label") or "",
            data.get("generic") if isinstance(data.get("generic"), dict) else None,
            data.get("clauses") if isinstance(data.get("clauses"), list) else None,
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    return jsonify(result)


@app.post("/api/agent/chat")
def agent_chat():
    data = request.get_json(silent=True) or {}
    try:
        flag = data.get("confirm")
        result = agent_reply(
            data.get("messages") or [],
            data.get("page") or "/",
            data.get("config") if isinstance(data.get("config"), dict) else None,
            data.get("threadId") if isinstance(data.get("threadId"), str) else None,
            flag if isinstance(flag, bool) else None,
            data.get("catalog") if isinstance(data.get("catalog"), dict) else None,
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    return jsonify(result)


@app.post("/api/rules/chat")
def rules_chat():
    data = request.get_json(silent=True) or {}
    try:
        result = discuss_rules(data.get("messages"), data.get("rules"), data.get("focus"))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    return jsonify(result)


@app.get("/api/modules")
def modules():
    return jsonify({"modules": module_catalog()})


@app.post("/api/calculate")
def calculate():
    data = request.get_json(silent=True) or {}
    upload_id = str(data.get("upload_id") or "")
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "请重新上传文件"}), 400
    specs = data.get("tables")
    if not isinstance(specs, list) or not specs:
        return jsonify({"error": "请至少添加一张计算表并对应表头"}), 400
    graph = load_graph()
    known = {item["id"] for item in graph.get("rules") or [] if isinstance(item, dict)}
    builtin = [spec for spec in specs if str(spec.get("product") or "") not in known]
    authored = [spec for spec in specs if str(spec.get("product") or "") in known]
    try:
        packs = []
        rows = []
        narrative = []
        exposure = {"currencies": []}
        if builtin:
            builtin_result = run_with_mappings(stored[0], builtin, calculation_rules())
            packs.extend(builtin_result.get("packs") or [])
            rows.extend(builtin_result.pop("rwa_rows", []))
            narrative.append(builtin_result.get("narrative") or "")
            exposure = builtin_result.get("exposure") or exposure
        if authored:
            authored_result = run_authored(stored[0], authored, graph)
            packs.extend(authored_result["packs"])
            rows.extend(authored_result["rwa_rows"])
            narrative.append("自建规则按结果公式计算。")
        result = {
            "mapping_source": "用户映射",
            "narrative": " ".join(part for part in narrative if part).strip(),
            "packs": packs,
            "pack_count": len(packs),
            "row_count": sum(pack.get("row_count", len(pack.get("rows") or [])) for pack in packs),
            "exposure": exposure,
        }
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"处理失败: {e}"}), 500
    original = stored[2] if len(stored) > 2 else "结果.xlsx"
    out_path = _result_path(stored[1], original)
    try:
        write_rwa_workbook(stored[0], out_path, rows)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"写出 Excel 失败: {e}"}), 500
    result["download_url"] = f"/api/result/{upload_id}"
    result["download_name"] = _result_name(original)
    exposure = result.get("exposure") or {"currencies": []}
    exposure["module"] = str(data.get("module") or "")
    EXPOSURES[upload_id] = exposure
    return jsonify(result)


@app.post("/api/compare")
def compare():
    upload_id = str(request.form.get("upload_id") or "")
    stored = UPLOADS.get(upload_id)
    exposure = EXPOSURES.get(upload_id)
    if stored is None or exposure is None:
        return jsonify({"error": "请先上传明细并完成计算"}), 400
    upload = request.files.get("file")
    if upload is None or upload.filename == "":
        return jsonify({"error": "请上传监管结果"}), 400
    name = upload.filename or ""
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "监管结果仅支持 .xlsx 或 .xlsm"}), 400
    reg_path = stored[1] / "regulatory.xlsx"
    upload.save(reg_path)
    try:
        module = str(request.form.get("module") or "")
        comparison = compare_module(module, exposure, reg_path)
        dest = stored[1] / "comparison.xlsx"
        write_comparison(dest, comparison)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"对比失败: {e}"}), 500
    comparison["download_url"] = f"/api/compare/{upload_id}"
    comparison["download_name"] = _compare_name(stored[2] if len(stored) > 2 else "结果.xlsx")
    return jsonify(comparison)


@app.get("/api/compare/<upload_id>")
def compare_file(upload_id: str):
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "请重新上传文件"}), 400
    path = stored[1] / "comparison.xlsx"
    if not path.is_file():
        return jsonify({"error": "还没有可下载的对比结果"}), 404
    name = _compare_name(stored[2] if len(stored) > 2 else "结果.xlsx")
    return send_file(path, as_attachment=True, download_name=name)


def _compare_name(filename: str) -> str:
    stem = Path(filename).stem or "结果"
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", stem).strip()[:80] or "结果"
    return f"{stem}_对比.xlsx"


def _preview_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        return str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).replace("\n", " ").strip()[:200]


def _looks_like_header(values: list) -> bool:
    texts = [item for item in values if isinstance(item, str) and item.strip() and not item.startswith("=")]
    return len(texts) >= 2


def _sheet_table(ws, get_column_letter, header_hint: int = 0) -> dict:
    header_row = 0
    headers: list = []
    stored = []
    total = 0
    hinted = header_hint if header_hint and header_hint > 0 else 0
    for index, row in enumerate(ws.iter_rows(max_col=40, values_only=True), start=1):
        values = list(row)
        if header_row == 0:
            if hinted and index < hinted:
                continue
            if hinted and index <= hinted + 3:
                if _looks_like_header(values):
                    header_row = index
                    headers = values
                continue
            if index <= 30 and _looks_like_header(values):
                header_row = index
                headers = values
            continue
        if not any(item is not None and str(item).strip() for item in values):
            continue
        total += 1
        if len(stored) < 300:
            stored.append((index, values))
        if total >= 5000:
            break
    columns = []
    width = 0
    for col, header in enumerate(headers):
        text = _preview_cell(header)
        if text:
            width = col + 1
    for col in range(width):
        header = _preview_cell(headers[col]) if col < len(headers) else ""
        if not header:
            continue
        samples = []
        for _excel_row, values in stored:
            if col >= len(values):
                continue
            text = _preview_cell(values[col])
            if text and text not in samples:
                samples.append(text[:40])
            if len(samples) >= 3:
                break
        columns.append({
            "letter": get_column_letter(col + 1),
            "header": header,
            "sample": (samples[0] if samples else "") or "—",
            "samples": samples,
        })
    rows = []
    for offset, packed in enumerate(stored, start=1):
        excel_row, values = packed
        item = {"id": "u" + str(offset), "excelRow": excel_row}
        for col in columns:
            index = 0
            letter = col["letter"]
            while get_column_letter(index + 1) != letter:
                index += 1
                if index > 50:
                    break
            item[letter] = _preview_cell(values[index]) if index < len(values) else ""
        rows.append(item)
    return {
        "headerRow": header_row,
        "columns": columns,
        "rows": rows,
        "rowCount": total,
        "truncated": total > len(rows),
    }


@app.post("/api/sheet-preview")
def sheet_preview():
    upload = request.files.get("file")
    if upload is None or upload.filename == "":
        return jsonify({"error": "请先选择要上传的 Excel 文件"}), 400
    name = upload.filename or ""
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "仅支持 .xlsx 或 .xlsm 格式"}), 400
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    work_dir = Path(tempfile.mkdtemp(prefix="rwa_preview_"))
    input_path = work_dir / "input.xlsx"
    kept = False
    try:
        upload.save(input_path)
        from agent.subagents.weight.paper.paper_calc import plan_workbook_sheets

        plan_book = load_workbook(input_path, read_only=True, data_only=True)
        try:
            planned, plan_note = plan_workbook_sheets(plan_book)
        except Exception as exc:
            planned, plan_note = [], "模型暂时没能判断表结构：" + str(exc)[:80]
        finally:
            plan_book.close()
        book = load_workbook(input_path, read_only=True, data_only=True)
        try:
            picked = None
            order = [item["sheet"] for item in planned] or list(book.sheetnames)
            hints = {item["sheet"]: item for item in planned}
            for sheet_name in order:
                if sheet_name not in book.sheetnames:
                    continue
                hint = hints.get(sheet_name) or {}
                table = _sheet_table(book[sheet_name], get_column_letter, int(hint.get("header_row") or 0))
                if table["columns"]:
                    picked = table
                    picked["sheet"] = sheet_name
                    picked["plan"] = planned
                    reason = str(hint.get("reason") or "").strip()
                    picked["planNote"] = plan_note or ("模型选定「" + sheet_name + "」作为明细。" + ((" " + reason) if reason else ""))
                    break
        finally:
            book.close()
        if picked is not None:
            upload_id = uuid.uuid4().hex
            UPLOADS[upload_id] = (input_path, work_dir, Path(name).name)
            picked["uploadId"] = upload_id
            picked["file"] = Path(name).name
            kept = True
    except Exception as exc:
        return jsonify({"error": "读取失败: " + str(exc)}), 500
    finally:
        if not kept:
            shutil.rmtree(work_dir, ignore_errors=True)
    if picked is None:
        return jsonify({"error": "这个工作簿里没有读到表头"}), 400
    return jsonify(picked)


@app.post("/api/sheet-map")
def sheet_map():
    """上传表头对应标准字段。交给模型，不在本地用关键词凑。"""
    data = request.get_json(silent=True) or {}
    headers = data.get("headers") or data.get("columns") or []
    fields = data.get("fields") or []
    if not isinstance(headers, list) or not isinstance(fields, list):
        return jsonify({"error": "headers 和 fields 须为数组"}), 400
    norm_headers = []
    for item in headers:
        if not isinstance(item, dict):
            continue
        letter = str(item.get("column") or item.get("letter") or "").strip().upper()
        title = str(item.get("header") or "").strip()
        if not re.fullmatch(r"[A-Z]+", letter) or not title:
            continue
        raw_samples = item.get("samples") if isinstance(item.get("samples"), list) else []
        if not raw_samples and item.get("sample"):
            raw_samples = [item.get("sample")]
        samples = []
        for sample in raw_samples:
            text = str(sample or "").strip()
            if text and text != "—" and text not in samples:
                samples.append(text[:40])
            if len(samples) >= 3:
                break
        norm_headers.append({"column": letter, "header": title[:80], "samples": samples})
    norm_fields = []
    for item in fields:
        if not isinstance(item, dict):
            continue
        fid = str(item.get("id") or "").strip()
        if not fid:
            continue
        raw_options = item.get("options") if isinstance(item.get("options"), list) else []
        options = [str(opt)[:48] for opt in raw_options[:30] if str(opt or "").strip()]
        norm_fields.append({
            "id": fid[:64],
            "label": str(item.get("label") or fid)[:40],
            "required": True,
            "hint": str(item.get("hint") or "")[:160],
            "sample": str(item.get("sample") or "")[:40],
            "options": options,
        })
    rows_in = data.get("rows") if isinstance(data.get("rows"), list) else []
    rows = []
    letters = [item["column"] for item in norm_headers]
    for item in rows_in[:40]:
        if not isinstance(item, dict):
            continue
        row = {}
        for letter in letters:
            text = str(item.get(letter) or "").strip()
            if text and text != "—":
                row[letter] = text[:40]
        if row:
            rows.append(row)
    table = {
        "file": str(data.get("file") or "")[:120],
        "sheet": str(data.get("sheet") or "")[:80],
        "header_row": data.get("header_row") or data.get("headerRow"),
        "row_count": data.get("row_count") or data.get("rowCount") or len(rows_in),
        "rows": rows,
        "note": "以下是用户上传的 Excel 内容。行数较多时只附上前面部分，列结构与整表一致。",
        "confirmed": str(data.get("decisions") or "").strip()[:400],
    }
    try:
        result = suggest_field_map(norm_headers, norm_fields, table)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    return jsonify(result)


@app.post("/api/sheet-clarify")
def sheet_clarify():
    """研判之后，模型列出拿不准的数据选择，等用户确认再映射。"""
    from openpyxl import load_workbook

    from agent.subagents.weight.paper.clarify import clarify_workbook

    data = request.get_json(silent=True) or {}
    upload_id = str(data.get("uploadId") or "")
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "这本工作簿的研判已经过期，请重新上传。"}), 400
    path = stored[0]
    messages = data.get("messages") if isinstance(data.get("messages"), list) else []
    plan = data.get("plan") if isinstance(data.get("plan"), list) else []
    note = str(data.get("note") or "")[:500]
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            result = clarify_workbook(book, plan, note, messages)
        finally:
            book.close()
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    if result.get("ready") and result.get("sheet"):
        from openpyxl.utils import get_column_letter

        sheet_name = result["sheet"]
        focus = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name in focus.sheetnames:
                table = _sheet_table(focus[sheet_name], get_column_letter, int(result.get("header_row") or 0))
                if table.get("columns"):
                    table["sheet"] = sheet_name
                    result["table"] = table
        finally:
            focus.close()
    return jsonify(result)


@app.post("/api/open-weights")
def open_weights():
    """规则树停住、没有固定权重时，让模型在规则允许的档次里选定，再按原规则重算。"""
    data = request.get_json(silent=True) or {}
    rows = data.get("rows") if isinstance(data.get("rows"), list) else []
    try:
        priced, note = price_open_facts(rows[:80])
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    return jsonify({"rows": priced, "note": note})


@app.post("/api/calc-trace")
def calc_trace():
    """把本次每一行的计算步骤写成 output/traces 下的 JSON，供以后倒查。"""
    from agent.subagents.weight.paper.calc_trace import save_trace_file

    data = request.get_json(silent=True) or {}
    rows = data.get("rows") if isinstance(data.get("rows"), list) else []
    path = save_trace_file(str(data.get("file") or "计算"), rows)
    return jsonify({"ok": True, "file": path.name, "rows": len([row for row in rows if isinstance(row, dict)])})


@app.post("/api/calc-export")
def calc_export():
    """下载时在源工作簿每一行后面写上风险暴露类型、权重、敞口和风险加权资产。"""
    data = request.get_json(silent=True) or {}
    lines = data.get("lines") if isinstance(data.get("lines"), list) else []
    source_lines = [line for line in lines if isinstance(line, dict) and line.get("excel_row") not in (None, "")]
    if source_lines:
        from agent.subagents.weight.paper.paper_calc import write_annotated_source
        from agent.subagents.weight.paper.working_paper import PAPER_PATH

        upload_id = str(data.get("uploadId") or "")
        stored = UPLOADS.get(upload_id) if upload_id else None
        if stored is not None:
            source, filename = stored[0], stored[2]
        elif PAPER_PATH.is_file():
            source, filename = PAPER_PATH, PAPER_PATH.name
        else:
            source, filename = None, ""
        if source is not None and Path(source).is_file():
            rows = []
            for line in source_lines:
                excel_row = line.get("excel_row")
                header_row = line.get("header_row")
                if isinstance(excel_row, str) and excel_row.isdigit():
                    excel_row = int(excel_row)
                if isinstance(header_row, str) and header_row.isdigit():
                    header_row = int(header_row)
                rows.append({
                    "sheet": line.get("sheet") or "",
                    "excel_row": excel_row,
                    "header_row": header_row,
                    "name": line.get("name") or "",
                    "exposure": line.get("exposure") or line.get("subclass") or "",
                    "weight_text": line.get("weight") or "",
                    "ead": line.get("ead") if isinstance(line.get("ead"), (int, float)) else None,
                    "rwa": line.get("rwa") if isinstance(line.get("rwa"), (int, float)) else None,
                    "trace": line.get("trace") or "",
                    "steps": line.get("steps") if isinstance(line.get("steps"), list) else [],
                })
            from agent.subagents.weight.paper.calc_trace import save_trace_file

            save_trace_file(filename or "计算", rows)
            buf = io.BytesIO()
            write_annotated_source({"process_rows": rows}, buf, Path(source))
            buf.seek(0)
            return send_file(
                buf,
                as_attachment=True,
                download_name=_result_name(filename),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "计算结果"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A5F4A")
    source_rows = any(isinstance(line, dict) and line.get("sheet") for line in lines)
    headers = (
        ["来源表", "行", "名称", "风险暴露类型", "权重", "敞口（亿元）", "风险加权资产（亿元）"]
        if source_rows
        else ["名称", "分类", "细分", "权重", "敞口（亿元）", "风险加权资产（亿元）", "公式"]
    )
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for line in lines:
        if not isinstance(line, dict):
            continue
        if source_rows:
            sheet.append([
                str(line.get("sheet") or "")[:40],
                line.get("excel_row") if isinstance(line.get("excel_row"), int) else None,
                str(line.get("name") or "")[:80],
                str(line.get("exposure") or line.get("subclass") or "")[:80],
                str(line.get("weight") or "")[:20],
                line.get("ead") if isinstance(line.get("ead"), (int, float)) else None,
                line.get("rwa") if isinstance(line.get("rwa"), (int, float)) else None,
            ])
        else:
            sheet.append([
                str(line.get("name") or "")[:80],
                str(line.get("category") or "")[:40],
                str(line.get("subclass") or "")[:40],
                str(line.get("weight") or "")[:20],
                line.get("ead") if isinstance(line.get("ead"), (int, float)) else None,
                line.get("rwa") if isinstance(line.get("rwa"), (int, float)) else None,
                str(line.get("formula") or "")[:180],
            ])
    sheet.append([])
    sheet.append(["合计敞口（亿元）", data.get("ead") if isinstance(data.get("ead"), (int, float)) else None])
    sheet.append(["合计风险加权资产（亿元）", data.get("rwa") if isinstance(data.get("rwa"), (int, float)) else None])
    for column, width in zip("ABCDEFG", (36, 28, 28, 14, 18, 26, 42)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    book.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="信用风险RWA计算结果.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/api/result/<upload_id>")
def result_file(upload_id: str):
    stored = UPLOADS.get(upload_id)
    if stored is None:
        return jsonify({"error": "请重新上传文件"}), 400
    original = stored[2] if len(stored) > 2 else "结果.xlsx"
    path = _result_path(stored[1], original)
    if not path.is_file():
        return jsonify({"error": "还没有可下载的计算结果"}), 404
    name = _result_name(original)
    return send_file(path, as_attachment=True, download_name=name)


# ---------------------------------------------------------------- 流动性 HQLA
# 判定引擎在 agent/subagents/hqla/，不进计算流水线的 STEPS。
# 任务数据的持久化在 Koa（MySQL），这里只做解析、计算和导出。


def _hqla_payload():
    return request.get_json(silent=True) or {}


@app.get("/api/hqla/list-meta")
def hqla_list_meta():
    from agent.subagents.hqla import list_meta

    return jsonify(list_meta())


@app.post("/api/hqla/inspect")
def hqla_inspect():
    from agent.subagents.hqla import inspect

    try:
        return jsonify(inspect(_hqla_payload()))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"解析失败：{exc}"}), 500


@app.post("/api/hqla/parse")
def hqla_parse():
    from agent.subagents.hqla import parse

    try:
        return jsonify(parse(_hqla_payload()))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"解析失败：{exc}"}), 500


@app.post("/api/hqla/compute")
def hqla_compute():
    from agent.subagents.hqla import compute

    try:
        return jsonify(compute(_hqla_payload()))
    except Exception as exc:
        return jsonify({"error": f"计算失败：{exc}"}), 500


@app.post("/api/hqla/sample")
def hqla_sample():
    from agent.subagents.hqla import sample_plan

    try:
        return jsonify(sample_plan(_hqla_payload()))
    except Exception as exc:
        return jsonify({"error": f"抽样失败：{exc}"}), 500


@app.post("/api/hqla/list-hits")
def hqla_list_hits():
    from agent.subagents.hqla import list_hits

    try:
        return jsonify(list_hits(_hqla_payload()))
    except Exception as exc:
        return jsonify({"error": f"匹配失败：{exc}"}), 500


@app.post("/api/hqla/export")
def hqla_export():
    from agent.subagents.hqla import export as export_hqla

    payload = _hqla_payload()
    try:
        data = export_hqla(payload)
    except Exception as exc:
        return jsonify({"error": f"导出失败：{exc}"}), 500
    task = payload.get("task") or {}
    base_date = (task.get("cfg") or {}).get("baseDate") or task.get("baseDate") or ""
    name = f"HQLA底稿_{task.get('name') or 'HQLA'}_{base_date}.xlsx".replace("/", "_")
    buf = io.BytesIO(data)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    loaded = ensure_library()
    port = int(os.environ.get("RWA_PORT", "8766"))
    print(f"已读取规则库 {len(loaded)} 条")
    print(f"市场风险 RWA 服务: http://127.0.0.1:{port}")
    app.run(host="127.0.0.1", port=port, debug=False)
