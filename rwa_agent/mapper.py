"""Ask the model to map formula fields onto the uploaded header row."""

from __future__ import annotations

import json
import os
import re
import urllib.error
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from rwa_agent.agent import PACK_FORMULA
from rwa_agent.fields import FIELD_HINT, FIELD_LABEL, OPTIONAL, PACK_TITLE, REQUIRED

QWEN_BASE = "https://dashscope.aliyuncs.com/compatible-mode/v1"
QWEN_MODEL = "qwen-plus"


def suggest_field_map(headers: list[dict], fields: list[dict], table: dict | None = None) -> dict[str, dict]:
    """把用户上传的 Excel 交给模型，说明要计算 RWA，由模型提取要用的列。"""
    if not headers:
        raise ValueError("这一行没有读到表头")
    clean_fields = [item for item in fields if isinstance(item, dict) and str(item.get("id") or "").strip()]
    if not clean_fields:
        raise ValueError("没有需要对应的字段")
    payload = {
        "task": "用户上传了这份 Excel，需要用它计算信用风险 RWA（风险加权资产）。请从表格内容里提取计算要用的列。",
        "workbook": table or {},
        "fields": clean_fields,
        "headers": headers,
    }
    raw = _complete(_field_system(), json.dumps(payload, ensure_ascii=False))
    parsed = _extract_json(raw)
    incoming = parsed.get("mapping") if isinstance(parsed.get("mapping"), dict) else parsed
    if not isinstance(incoming, dict):
        incoming = {}
    allowed = {str(item.get("column") or "") for item in headers}
    used: set[str] = set()
    mapped: dict[str, dict] = {}
    for field in clean_fields:
        fid = str(field.get("id") or "").strip()
        slot = incoming.get(fid)
        letter = ""
        const = ""
        reason = ""
        if isinstance(slot, dict):
            letter = str(slot.get("column") or slot.get("letter") or "").strip().upper()
            const = str(slot.get("const") or slot.get("value") or "").strip()
            reason = str(slot.get("reason") or "").strip()[:80]
        elif isinstance(slot, str):
            token = slot.strip()
            if re.fullmatch(r"[A-Za-z]+", token):
                letter = token.upper()
            else:
                const = token
        if re.fullmatch(r"[A-Z]+", letter) and letter in allowed and letter not in used:
            mapped[fid] = {"column": letter, "reason": reason or ("模型对应到 " + letter)}
            used.add(letter)
            continue
        if not const:
            const = str(field.get("sample") or "").strip() or "未提供"
            reason = reason or "表里没有能对上的列，先用字段示例作常量"
        mapped[fid] = {"const": const[:80], "reason": reason or "表里没有这一列，模型填了常量"}
    extract = str(parsed.get("extract") or "").strip()[:240]
    return {"mapping": mapped, "extract": extract}


def _field_system() -> str:
    return (
        "用户上传了一份 Excel，用途是计算信用风险 RWA（风险加权资产）。"
        "fields 里的 id 必须原样使用，不要改名，也不要漏掉任何一个字段。"
        "你要阅读这份表格，把每一个字段对到表头，或在对不上时填常量。"
        "只输出一个 JSON 对象，不要解释。"
        "格式是 {\"extract\": \"一两句：从表里提取了哪些计算要用的内容\", "
        "\"mapping\": {\"字段id\": {\"column\": \"列字母\", \"reason\": \"一句为什么这列能用于计算\"}}}。"
        "对得上表头时写 column。一个表头只能分给一个字段，不要用合同号、序号、分区去凑。"
        "金额（balance、single_exposure、账面余额、本金、敞口）对金额列，不要对到权重或 RWA。"
        "对不上表头时不要留空，改写 {\"const\": \"常量\", \"reason\": \"表里为什么没有这一列\"}。"
        "常量优先用这张表能看出来的取值，看不出来再用该字段的 sample。"
        "常见对应：产品类型或业务品种对 product_type、instrument_type；"
        "发行人类型或客户类型对 issuer_type、counterparty_type；"
        "评级对 issuer_rating；清偿顺序对 seniority；是否上市对 listed；"
        "期限对 maturity；余额对 balance；减值对 impairment；五级分类或违约对 default_status、is_defaulted。"
        "options 里是「中文=code」。选列时不要改单元格；填常量时用表里的原文或 sample。"
    )


def suggest_mappings(path: str | Path, sheet_name: str, header_row: int, products: list[str]) -> dict[str, dict[str, str]]:
    wanted = [product for product in products if product in REQUIRED]
    if not wanted:
        raise ValueError("请先选择计算方式")
    headers = _header_samples(path, sheet_name, header_row)
    if not headers:
        raise ValueError("这一行没有读到表头")
    raw = _complete(_system(), _user(wanted, headers))
    parsed = _extract_json(raw)
    blob = parsed.get("mappings") if isinstance(parsed.get("mappings"), dict) else parsed
    allowed = {item["column"] for item in headers}
    result: dict[str, dict[str, str]] = {}
    for product in wanted:
        incoming = blob.get(product) if isinstance(blob, dict) else None
        if not isinstance(incoming, dict):
            result[product] = {}
            continue
        used: set[str] = set()
        columns: dict[str, str] = {}
        by_letter = {item["column"]: item for item in headers}
        fields = REQUIRED[product] + OPTIONAL[product]
        for field in fields:
            letter = str(incoming.get(field) or "").strip().upper()
            if field == "position_local":
                letter = _prefer_ead(headers, letter)
            if field in ("issuer_type", "counterparty_type"):
                letter = _prefer_new(headers, letter)
            if not re.fullmatch(r"[A-Z]+", letter) or letter not in allowed or letter in used:
                continue
            if not _accept(field, by_letter[letter]):
                continue
            columns[field] = letter
            used.add(letter)
        result[product] = columns
    return result


def _accept(field: str, column: dict) -> bool:
    samples = [str(item) for item in column.get("samples") or []]
    title = str(column.get("header") or "")
    upper = title.upper()
    if field == "issuer_name" and any(_is_code(sample) for sample in samples):
        return False
    if field in ("rating_sp", "rating_fitch") and samples and all(_is_moodys(sample) for sample in samples):
        return False
    if field.startswith("far_") and not any(token in upper for token in ("FAR", "远端", "LEG")):
        return False
    if field in ("position_local", "mv_mop", "nop_mop") and ("RWA" in upper or upper.startswith("RW")):
        return False
    return True


def _prefer_new(headers: list[dict], letter: str) -> str:
    chosen = next((item for item in headers if item["column"] == letter), None)
    if chosen is None or "ORIGINAL" not in chosen["header"].upper():
        return letter
    stem = chosen["header"].upper().replace("_ORIGINAL", "").replace("ORIGINAL", "")
    for item in headers:
        title = item["header"].upper()
        if "NEW" in title and stem and stem in title.replace("_NEW", "").replace("NEW", ""):
            return item["column"]
    return letter


def _prefer_ead(headers: list[dict], letter: str) -> str:
    chosen = next((item for item in headers if item["column"] == letter), None)
    if chosen is None or "PRE" not in chosen["header"].upper():
        return letter
    for item in headers:
        title = item["header"].upper().replace(" ", "")
        if title.endswith("EAD") and "PRE" not in title and "PRO" not in title and "RWA" not in title:
            return item["column"]
    return letter


def _is_code(sample: str) -> bool:
    text = sample.strip()
    return bool(re.match(r"^(CHN|HKG|S)\d", text) or re.match(r"^[A-Z]{2,4}-[\w-]+$", text))


def _is_moodys(sample: str) -> bool:
    text = sample.strip()
    return bool(re.fullmatch(r"(Aaa|Aa[1-3]|A[1-3]|Baa[1-3]|Ba[1-3]|B[1-3]|Caa[1-3]|Ca|C)", text))


def _header_samples(path: str | Path, sheet_name: str, header_row: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"没有工作表「{sheet_name}」")
        ws = wb[sheet_name]
        header: list = []
        samples: list[list] = []
        for index, row in enumerate(ws.iter_rows(max_col=45, values_only=True), start=1):
            values = list(row)
            if index < header_row:
                continue
            if index == header_row:
                header = values
                continue
            if any(cell is not None and str(cell).strip() for cell in values):
                samples.append(values)
            if len(samples) >= 3:
                break
    finally:
        wb.close()
    found = []
    for col, cell in enumerate(header):
        title = _one_line(cell)
        if not title:
            continue
        examples = []
        for sample in samples:
            if col >= len(sample):
                continue
            text = _one_line(sample[col])
            if text and text not in examples:
                examples.append(text[:40])
        found.append({"column": get_column_letter(col + 1), "header": title[:80], "samples": examples[:3]})
    return found


def _user(products: list[str], headers: list[dict]) -> str:
    packs = []
    for product in products:
        fields = []
        for field in REQUIRED[product] + OPTIONAL[product]:
            fields.append(
                {
                    "id": field,
                    "label": FIELD_LABEL.get(field, field),
                    "required": field in REQUIRED[product],
                    "hint": FIELD_HINT.get(field, ""),
                }
            )
        packs.append(
            {
                "id": product,
                "title": PACK_TITLE.get(product, product),
                "formula": PACK_FORMULA.get(product, ""),
                "fields": fields,
            }
        )
    return json.dumps({"formulas": packs, "headers": headers}, ensure_ascii=False)


def _system() -> str:
    return (
        "你负责把资本公式的字段映射到上传表的表头。只输出一个 JSON 对象，不要解释。"
        "格式是 {\"mappings\": {\"公式id\": {\"字段id\": \"列字母\"}}}。"
        "一个表头只能分给一个字段。含义对不上就不要填这个字段，不要用序号、分区、item 去凑。"
        "发行人要的是名称文本，客户号和合约号不要填到发行人。"
        "S&P 评级只接受 AAA、AA+、BBB- 这种标普符号。示例若是 A1、Aa1、Aaa，那是穆迪，不要填到 rating_sp。"
        "对手类型、发行人类型优先看主体类型，例如 BIS_ENTITY_TYPE。"
        "币种看 CURRENCY。持仓或敞口金额看 EAD，不要看 RWA、风险权重。"
        "剩余期限看 RESIDUAL_MATURITY。交易编号看合约号 CONTRACT_REFERENCE。"
        "表里没有远端两腿金额时，far_ccy、far_amt 留空。"
    )


def suggest_complete_map(headers: list[dict], fields: list[dict], table: dict | None = None) -> dict[str, dict]:
    """每个分类字段都要有结论：对到列、整表常量，或底稿没有。同一列可以对应多个字段。"""
    if not headers:
        raise ValueError("这一行没有读到表头")
    clean_fields = [item for item in fields if isinstance(item, dict) and str(item.get("id") or "").strip()]
    if not clean_fields:
        raise ValueError("没有需要对应的字段")
    payload = {
        "task": "这是明天演示用的信用风险底稿。请把每一个字段都对应上，然后才能计算 RWA。",
        "workbook": table or {},
        "fields": clean_fields,
        "headers": headers,
    }
    raw = _complete(_complete_system(), json.dumps(payload, ensure_ascii=False), timeout=120)
    parsed = _extract_json(raw)
    incoming = parsed.get("mapping") if isinstance(parsed.get("mapping"), dict) else {}
    allowed = {str(item.get("column") or "") for item in headers}
    found: dict[str, dict] = {}
    for field in clean_fields:
        fid = str(field.get("id") or "").strip()
        slot = incoming.get(fid) if isinstance(incoming, dict) else None
        found[fid] = _normalize_slot(slot, allowed)
    return {"mapping": found, "extract": str(parsed.get("extract") or "").strip()[:240]}


def _complete_system() -> str:
    return (
        "用户给出信用风险底稿的一张表，要计算 RWA（风险加权资产）。"
        "字段 id 来自权重法风险暴露分类 JSON，必须原样使用。"
        "每一个字段都要在 mapping 里出现一次，不能漏。"
        "只输出一个 JSON 对象。"
        "格式是 {\"extract\":\"一句\",\"mapping\":{\"字段id\":{...}}}。"
        "对得上某一列时写 {\"column\":\"列字母\",\"reason\":\"为什么这列就是这个字段\"}。"
        "整张表都是同一个取值时写 {\"const\":\"选项code\",\"reason\":\"为什么整表如此\"}，const 必须是该字段 options 里的 code。"
        "底稿确实没有、也不能从这张表的产品性质判断时写 {\"missing\":true,\"reason\":\"底稿没有\"}。"
        "同一列可以分给多个字段。不要编造不存在的列字母。"
        "single_exposure 只能对账面余额：Value Balance(MOP)、O/S BAL_MOP、Equiv to MOP、信貸敞口_MOP、剩余可使用额度（澳门元）。"
        "抵押品价值、Collateral Value、利息、RWA、风险权重、CCF、合同号、客户号都不是 single_exposure。"
        "Cust Type、INDUSTRY、客户类型对 counterparty_type。存放同业整表是商业银行，不要写成主权。"
        "信用卡整表是自然人，并且 is_credit_card 为 yes。贷款、按揭、透支、额度、信用卡的 fi_bank_applicable 为 not_applicable。"
        "原始期限对 Rest Period、Original Period、BY PERIOD (DAYS)，这些数是天数。不要对到期日。"
        "五级分类、DPD、逾期天数对 is_defaulted。底稿没有的财务和资本条件写 missing，不要整表填 yes。"
        "抵押品描述对 property_uses 和 has_property_collateral。"
        "不要把风险权重、RWA、合同号、客户号当成分类字段。"
    )


def _normalize_slot(slot: object, allowed: set[str]) -> dict:
    if isinstance(slot, str) and re.fullmatch(r"[A-Za-z]+", slot.strip()):
        slot = {"column": slot.strip().upper()}
    if not isinstance(slot, dict):
        return {"status": "missing", "reason": "模型没有给出这一项"}
    if slot.get("missing") is True or str(slot.get("status") or "") == "missing":
        return {"status": "missing", "reason": str(slot.get("reason") or "底稿没有")[:80]}
    letter = str(slot.get("column") or slot.get("letter") or "").strip().upper()
    const = slot.get("const")
    reason = str(slot.get("reason") or "").strip()[:80]
    if letter and letter in allowed:
        return {"status": "column", "column": letter, "reason": reason or ("模型对到 " + letter)}
    if const is not None and str(const).strip() and str(const).strip().lower() not in {"null", "none", "unknown"}:
        return {"status": "const", "const": const if isinstance(const, (int, float)) else str(const).strip(), "reason": reason or "模型按整表取值"}
    return {"status": "missing", "reason": reason or "底稿没有"}


def _complete(system: str, user: str, timeout: int = 90) -> str:
    key = (
        os.environ.get("DASHSCOPE_API_KEY", "").strip()
        or os.environ.get("QWEN_API_KEY", "").strip()
    )
    if not key:
        raise RuntimeError("未配置通义千问 API Key")
    base = (os.environ.get("QWEN_BASE_URL") or QWEN_BASE).strip()
    model = (os.environ.get("QWEN_MODEL") or QWEN_MODEL).strip()
    body = json.dumps(
        {
            "model": model,
            "temperature": 0,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        base.rstrip("/") + "/chat/completions",
        data=body,
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            payload = json.loads(resp.read().decode())
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode(errors="replace")[:300]
        raise RuntimeError(f"模型请求失败：{exc.code} {detail}") from exc
    return payload["choices"][0]["message"]["content"]


def _extract_json(text: str) -> dict:
    raw = text.strip()
    fenced = re.search(r"```(?:json)?\s*(\{.*\})\s*```", raw, re.DOTALL)
    if fenced:
        raw = fenced.group(1)
    start = raw.find("{")
    end = raw.rfind("}")
    if start == -1 or end <= start:
        raise RuntimeError("模型没有返回映射")
    return json.loads(raw[start : end + 1])


def _one_line(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text
