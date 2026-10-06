"""Choose a formula pack from the uploaded sheet, then calculate. The model may only remap headers."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from rwa_agent.fields import (
    FIELD_HINT,
    FIELD_LABEL,
    OPTIONAL,
    PACK_TITLE,
    REQUIRED,
    classify,
    map_header_row,
    norm_header,
)
from rwa_agent.formulas import (
    calc_ccr,
    calc_mr_fx,
    calc_mr_irs,
    calc_mr_repo,
    calc_repo_credit,
    fx_portfolio,
)
from rwa_agent.rules import resolve_rules

PREVIEW_ROWS = 40
TYPE_DISTINCT_MAX = 48
DETAIL_ROW_LIMIT = 40
# Longer names are tested first. Rank below 15 is a default type column.
_TYPE_RANK = (
    ("bisproductcategory", 0),
    ("natureofitem", 1),
    ("assetclassnew", 2),
    ("assetclassoriginal", 3),
    ("bisentitytypenew", 4),
    ("bisentitytypeoriginal", 5),
    ("bstype2", 6),
    ("bstype", 7),
    ("tradingbooktype", 8),
    ("itemafcrm", 14),
    ("item", 12),
)

CALCULATORS = {
    "repo_credit": calc_repo_credit,
    "ccr_fx": calc_ccr,
    "mr_repo": calc_mr_repo,
    "mr_irs": calc_mr_irs,
    "mr_fx": calc_mr_fx,
}

PACK_FORMULA = {
    "repo_credit": "RWA = 折美元持仓 × 发行人风险权重；资本 = RWA × 8%。发行人是中华人民共和国财政部时权重为 0。",
    "ccr_fx": "信用暴露 = 远端两腿折现后的净流入（小于 0 则记 0）+ 流入金额 × 外汇附加因子；RWA = 信用暴露 × 对手风险权重。",
    "mr_repo": "资本 = 市值 × 特定风险因子 + 市值 × 期限档权重；RWA = 资本 × 12.5。",
    "mr_irs": "资本 = |浮动腿现值 × 期限档权重 − 固定腿现值 × 期限档权重| × 2；RWA = 资本 × 12.5。",
    "mr_fx": "单币种资本 = 净敞口 × 8%；单币种 RWA = 净敞口。全行另取多头合计与空头合计的较大者。",
}


def field_catalog() -> list[dict]:
    catalog = []
    for product, title in PACK_TITLE.items():
        fields = [
            {"id": field, "label": FIELD_LABEL[field], "required": True, "hint": FIELD_HINT.get(field, "")}
            for field in REQUIRED[product]
        ]
        fields.extend(
            {
                "id": field,
                "label": FIELD_LABEL[field],
                "required": False,
                "hint": FIELD_HINT.get(field, ""),
            }
            for field in OPTIONAL[product]
        )
        catalog.append(
            {
                "id": product,
                "title": title,
                "formula": PACK_FORMULA[product],
                "fields": fields,
            }
        )
    return catalog


def inspect_workbook(path: str | Path) -> dict:
    scanned = _scan_workbook(path)
    sheets = [{"name": item["name"], "rows": item["rows"]} for item in scanned]
    previews = []
    for item in scanned:
        preview = _preview_sheet(item)
        if not preview["rows"] and not item["type_columns"]:
            continue
        preview["profile_header_row"] = item["profile_header_row"]
        preview["type_columns"] = item["type_columns"]
        preview["data_rows"] = item["data_rows"]
        previews.append(preview)
    return {
        "sheets": previews,
        "suggestions": [_suggestion(table) for table in _find_tables(sheets)],
        "catalog": field_catalog(),
    }


def list_type_columns(path: str | Path, sheet_name: str, header_row: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"没有工作表「{sheet_name}」")
        return _type_columns_of(wb[sheet_name], header_row)
    finally:
        wb.close()


def run_with_mappings(path: str | Path, specs: list[dict], rules: dict | None = None) -> dict:
    resolved = resolve_rules(rules)
    names = [str(spec.get("sheet") or "") for spec in specs]
    loaded = {sheet["name"]: sheet for sheet in _load_sheets(path, names, max_row=None)}
    tables = [_table_from_spec(loaded, spec) for spec in specs]
    packs = [_run_table(table, resolved) for table in tables]
    rwa_rows = []
    for pack in packs:
        rwa_rows.extend(pack.pop("_rwa_rows", []))
    return {
        "mapping_source": "用户映射",
        "narrative": _narrative(packs),
        "packs": packs,
        "pack_count": len(packs),
        "row_count": sum(pack.get("row_count", len(pack["rows"])) for pack in packs),
        "rwa_rows": rwa_rows,
    }


def _scan_workbook(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    scanned = []
    try:
        for name in wb.sheetnames:
            scanned.append(_scan_sheet(name, wb[name]))
    finally:
        wb.close()
    return scanned


def _scan_sheet(name: str, ws) -> dict:
    preview: list[list] = []
    header_idx = None
    header: list = []
    counts: list[Counter | None] = []
    data_rows = 0
    for index, row in enumerate(ws.iter_rows(max_col=45, values_only=True)):
        values = list(row)
        if index < PREVIEW_ROWS:
            preview.append(values)
        if header_idx is None:
            if index < 20 and _headerish(values):
                header_idx = index
                header = values
                counts = [Counter() for _ in header]
            continue
        if _is_blank(values) or _is_total(values):
            continue
        data_rows += 1
        if data_rows > 80000:
            break
        for col, counter in enumerate(counts):
            if counter is None or col >= len(values):
                continue
            token = _type_token(values[col])
            if token is None:
                continue
            counter[token] += 1
            if len(counter) > TYPE_DISTINCT_MAX:
                counts[col] = None
    return {
        "name": name,
        "rows": preview,
        "profile_header_row": None if header_idx is None else header_idx + 1,
        "type_columns": _type_column_view(header, counts),
        "data_rows": data_rows,
    }


def _type_columns_of(ws, header_row: int) -> list[dict]:
    if header_row < 1:
        raise ValueError("表头行无效")
    header: list = []
    counts: list[Counter | None] = []
    data_rows = 0
    for index, row in enumerate(ws.iter_rows(max_col=45, values_only=True), start=1):
        values = list(row)
        if index < header_row:
            continue
        if index == header_row:
            header = values
            counts = [Counter() for _ in header]
            continue
        if _is_blank(values) or _is_total(values):
            continue
        data_rows += 1
        if data_rows > 80000:
            break
        for col, counter in enumerate(counts):
            if counter is None or col >= len(values):
                continue
            token = _type_token(values[col])
            if token is None:
                continue
            counter[token] += 1
            if len(counter) > TYPE_DISTINCT_MAX:
                counts[col] = None
    return _type_column_view(header, counts)


def _type_column_view(header: list, counts: list[Counter | None]) -> list[dict]:
    found = []
    for col, counter in enumerate(counts):
        if not isinstance(counter, Counter) or len(counter) < 2:
            continue
        title = _one_line(header[col]) if col < len(header) else ""
        if not title:
            continue
        found.append(
            {
                "column": get_column_letter(col + 1),
                "header": title[:80],
                "rank": _type_rank(title),
                "values": [{"value": value, "count": count} for value, count in counter.most_common()],
            }
        )
    found.sort(key=lambda item: (item["rank"], item["column"]))
    return found


def _type_rank(header: str) -> int:
    text = norm_header(header)
    best = 100
    for needle, rank in sorted(_TYPE_RANK, key=lambda item: len(item[0]), reverse=True):
        if text == needle or (len(needle) > 4 and needle in text):
            return min(best, rank)
    return best


def _headerish(row: list) -> bool:
    texts = 0
    for cell in row:
        if isinstance(cell, str) and cell.strip() and not cell.startswith("="):
            texts += 1
    return texts >= 4


def _type_token(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        if value.startswith("="):
            return None
        text = value.strip()
        return text or None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9 and abs(value) < 1e12:
            return str(int(round(value)))
        return None
    text = str(value).strip()
    return text or None


def _preview_sheet(sheet: dict) -> dict:
    rows_out = []
    for index, row in enumerate(sheet["rows"][:PREVIEW_ROWS], start=1):
        cells = []
        for col, value in enumerate(row[:40]):
            text = _cell_text(value)
            if text:
                cells.append({"column": get_column_letter(col + 1), "text": text})
        if cells:
            rows_out.append({"row": index, "cells": cells})
    return {"name": sheet["name"], "rows": rows_out}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        if value.startswith("="):
            return ""
        return _one_line(value)[:80]
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-6:
            return f"{int(round(value)):,}"
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    if isinstance(value, int):
        return f"{value:,}"
    return _one_line(str(value))[:80]


def _suggestion(table: dict) -> dict:
    product = table["product"]
    wanted = REQUIRED[product] + OPTIONAL[product]
    columns = {}
    for field in wanted:
        source = field
        if field not in table["mapping"] and field == "issuer_category":
            source = "issuer_type"
        if field not in table["mapping"] and field == "counterparty_type":
            source = "issuer_type"
        col = table["mapping"].get(source)
        if col is None and source != field:
            col = table["mapping"].get(field)
        if col is not None:
            columns[field] = get_column_letter(col + 1)
    return {
        "sheet": table["sheet"],
        "header_row": table["header_row"],
        "product": product,
        "columns": columns,
    }


def _table_from_spec(loaded: dict[str, dict], spec: dict) -> dict:
    name = str(spec.get("sheet") or "")
    sheet = loaded.get(name)
    if sheet is None:
        raise ValueError(f"没有工作表「{name}」")
    product = spec.get("product")
    if product not in CALCULATORS:
        raise ValueError("请选择一种计算方式")
    try:
        header_row = int(spec.get("header_row"))
    except (TypeError, ValueError):
        raise ValueError(f"「{name}」的表头行无效") from None
    if header_row < 1 or header_row > len(sheet["rows"]):
        raise ValueError(f"「{name}」的表头行无效")
    mapping: dict[str, int] = {}
    for field, letter in (spec.get("columns") or {}).items():
        if field not in FIELD_LABEL or not letter:
            continue
        token = str(letter).strip().upper()
        if not re.fullmatch(r"[A-Z]+", token):
            continue
        mapping[field] = _col_index(token)
    type_index = None
    allowed: set[str] | None = None
    type_values = [str(value) for value in (spec.get("type_values") or []) if str(value).strip()]
    type_letter = str(spec.get("type_column") or "").strip().upper()
    if type_letter or type_values:
        if not re.fullmatch(r"[A-Z]+", type_letter) or not type_values:
            raise ValueError(f"「{name}」请先选定类型列，再选择要计算的类型")
        type_index = _col_index(type_letter)
        allowed = set(type_values)
    data_rows = []
    for offset, raw in enumerate(sheet["rows"][header_row:]):
        if _is_blank(raw):
            if data_rows and allowed is None:
                break
            continue
        if _is_total(raw):
            if allowed is None:
                break
            continue
        if allowed is not None:
            cell = raw[type_index] if type_index is not None and type_index < len(raw) else None
            if _type_token(cell) not in allowed:
                continue
        record = _record(raw, mapping, header_row + 1 + offset)
        if _has_payload(record, product):
            data_rows.append(record)
    book = "banking" if product == "repo_credit" else "trading" if str(product).startswith("mr_") else None
    return {
        "sheet": name,
        "header_row": header_row,
        "product": product,
        "book": book,
        "mapping": mapping,
        "header": sheet["rows"][header_row - 1],
        "rows": data_rows,
        "type_column": type_letter,
        "type_values": type_values,
    }


def run_agent(path: str | Path) -> dict:
    sheets = _load_sheets(path)
    tables = _find_tables(sheets)
    mapping_source = "字段别名"
    packs = [_run_table(table, resolve_rules(None)) for table in tables]
    for pack in packs:
        pack.pop("_rwa_rows", None)
    return {
        "mapping_source": mapping_source,
        "narrative": _narrative(packs),
        "packs": packs,
        "pack_count": len(packs),
        "row_count": sum(pack.get("row_count", len(pack["rows"])) for pack in packs),
    }


def _load_sheets(path: str | Path, names: list[str] | None = None, max_row: int | None = 200) -> list[dict]:
    wanted = None if names is None else {name for name in names if name}
    wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    sheets = []
    try:
        for name in wb.sheetnames:
            if wanted is not None and name not in wanted:
                continue
            rows = []
            for index, row in enumerate(wb[name].iter_rows(max_col=45, values_only=True), start=1):
                if max_row is not None and index > max_row:
                    break
                rows.append(list(row))
            sheets.append({"name": name, "rows": rows})
    finally:
        wb.close()
    return sheets


def _find_tables(sheets: list[dict]) -> list[dict]:
    found = []
    for sheet in sheets:
        rows = sheet["rows"]
        headers = [i for i, row in enumerate(rows) if len(map_header_row(row)) >= 3]
        for n, index in enumerate(headers):
            mapping = map_header_row(rows[index])
            product, book = classify(set(mapping), sheet["name"])
            if not product:
                continue
            end = headers[n + 1] if n + 1 < len(headers) else len(rows)
            data_rows = []
            for offset, raw in enumerate(rows[index + 1 : end]):
                if _is_blank(raw):
                    if data_rows:
                        break
                    continue
                if _is_total(raw):
                    break
                record = _record(raw, mapping, index + 2 + offset)
                if _has_payload(record, product):
                    data_rows.append(record)
            if data_rows:
                found.append(
                    {
                        "sheet": sheet["name"],
                        "header_row": index + 1,
                        "product": product,
                        "book": book,
                        "mapping": mapping,
                        "header": rows[index],
                        "rows": data_rows,
                    }
                )
    return found


def _record(raw: list, mapping: dict[str, int], excel_row: int) -> dict:
    record: dict = {"_row": excel_row}
    for field, col in mapping.items():
        value = raw[col] if col < len(raw) else None
        if isinstance(value, str) and value.startswith("="):
            value = None
        record[field] = value
    if record.get("issuer_category") in (None, "") and record.get("issuer_type"):
        record["issuer_category"] = record.get("issuer_type")
    if record.get("counterparty_type") in (None, "") and record.get("issuer_type"):
        record["counterparty_type"] = record.get("issuer_type")
    return record


def _has_payload(record: dict, product: str) -> bool:
    for field in REQUIRED[product]:
        value = record.get(field)
        if value not in (None, ""):
            return True
    return False


def _is_blank(row: list) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def _is_total(row: list) -> bool:
    for cell in row:
        if isinstance(cell, str) and ("合计" in cell or cell.strip().lower().startswith("total")):
            return True
    return False


def _run_table(table: dict, rules: dict) -> dict:
    product = table["product"]
    calc = CALCULATORS[product]
    rows = []
    for record in table["rows"]:
        rows.append(_run_row(product, record, calc, rules))
    currency = next((row["currency"] for row in rows if row.get("currency")), None)
    ok_rows = [row for row in rows if row["status"] == "ok"]
    title = PACK_TITLE[product]
    type_values = table.get("type_values") or []
    if type_values:
        shown = "、".join(type_values[:6])
        if len(type_values) > 6:
            shown += f" 等 {len(type_values)} 个类型"
        title = f"{title} · {shown}"
    notes = _pack_notes(product, table["book"])
    if type_values:
        notes.insert(0, "只计算类型列里取值为「" + "、".join(type_values[:6]) + "」的明细。")
    pack = {
        "id": product,
        "title": title,
        "sheet": table["sheet"],
        "book": table["book"],
        "header_row": table["header_row"],
        "formula": PACK_FORMULA[product],
        "currency": currency,
        "mapping": _mapping_view(table),
        "type_values": type_values,
        "row_count": len(rows),
        "rows": rows[:DETAIL_ROW_LIMIT],
        "rows_hidden": max(0, len(rows) - DETAIL_ROW_LIMIT),
        "total_rwa": sum(row["rwa"] for row in ok_rows),
        "total_capital": sum(row["capital"] for row in ok_rows),
        "notes": notes,
        "_rwa_rows": [
            {
                "sheet": table["sheet"],
                "header_row": table["header_row"],
                "excel_row": row["excel_row"],
                "rwa": row["rwa"],
            }
            for row in ok_rows
            if isinstance(row.get("rwa"), (int, float))
        ],
    }
    if product == "mr_fx" and ok_rows:
        pack["portfolio"] = fx_portfolio(ok_rows, rules["mr_fx"]["capital_ratio"])
    return pack


def _run_row(product: str, record: dict, calc, rules: dict) -> dict:
    missing = [field for field in REQUIRED[product] if record.get(field) in (None, "")]
    trade_id = record.get("trade_id") or f"第{record['_row']}行"
    base = {
        "id": str(trade_id),
        "excel_row": record["_row"],
        "missing": [FIELD_LABEL.get(field, field) for field in missing],
    }
    if missing:
        base.update(
            status="missing",
            rwa=None,
            capital=None,
            currency=None,
            summary="缺 " + "、".join(base["missing"]) + "，这一笔没有计算",
            steps=[],
        )
        return base
    try:
        result = calc(record, rules)
    except (TypeError, ValueError) as exc:
        base.update(
            status="missing",
            rwa=None,
            capital=None,
            currency=None,
            summary=f"数值无法计算：{exc}",
            steps=[],
        )
        return base
    if not result.get("ok"):
        base.update(
            status="missing",
            rwa=None,
            capital=None,
            currency=None,
            summary=result.get("message") or "无法计算",
            steps=result.get("steps") or [],
        )
        return base
    base.update(
        status="ok",
        rwa=result["rwa"],
        capital=result["capital"],
        currency=result["currency"],
        summary=result["summary"],
        steps=result["steps"],
    )
    if "ladder_capital" in result:
        base["ladder_capital"] = result["ladder_capital"]
    if "nop" in result:
        base["nop"] = result["nop"]
        base["short"] = result["short"]
    return base


def _mapping_view(table: dict) -> list[dict]:
    header = table["header"]
    view = []
    for field, col in sorted(table["mapping"].items(), key=lambda item: item[1]):
        title = header[col] if col < len(header) else ""
        view.append(
            {
                "column": get_column_letter(col + 1),
                "header": _one_line(title),
                "field": field,
                "label": FIELD_LABEL.get(field, field),
            }
        )
    return view


def _one_line(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value).strip()


def _pack_notes(product: str, book: str | None) -> list[str]:
    notes = []
    if product == "repo_credit":
        where = "银行账簿" if book != "trading" else "交易账簿"
        notes.append(f"表被识别为{where}正回购，债券本体走发行人信用风险。")
        notes.append("通告 025 其余 7 条需要次级债、担保人、原到期日或出资比例。模板里没有这些字段，对应特殊权重不启用。")
    elif product == "ccr_fx":
        notes.append("只读远端两腿。近端是否已结算没有进入公式。")
        notes.append("银行对手按「银行-一般」和 Fitch 评级。通告第 5 条的 30% 需要原到期日，模板没有该字段，A 档仍用标准法 50%。")
    elif product == "mr_repo":
        notes.append("交易账簿质押债。特定风险因子和期限档由发行人类别、剩余期限查参照表，不采用模板里填写的权重。")
        notes.append("同一债券若已在银行账簿计过发行人信用风险，这里不再相加。")
    elif product == "mr_irs":
        notes.append("名义本金不参与。资本用的是示例里的两腿差额绝对值乘 2。")
    elif product == "mr_fx":
        notes.append("按行的 RWA 等于该币种净敞口。全行口径另取多头合计与空头合计的较大者。")
    return notes


def _narrative(packs: list[dict]) -> str:
    if not packs:
        return "没有识别到回购或互换计算表。表头需要能对上持仓、远端金额、市值或互换现值这些字段。"
    lines = [f"识别到 {len(packs)} 张计算表。金额由公式写出，上传表里的权重和 RWA 不读入。"]
    for pack in packs:
        ok = sum(1 for row in pack["rows"] if row["status"] == "ok")
        bad = len(pack["rows"]) - ok
        money = _money(pack["total_rwa"], pack.get("currency"))
        lines.append(f"{pack['sheet']}：{pack['title']}，{ok} 笔已计算，RWA {money}。")
        if bad:
            lines.append(f"其中 {bad} 笔缺字段，没有出数。")
    lines.append("信用风险的美元结果和市场风险的澳门元结果分列，不加总。")
    return "".join(lines)


def write_rwa_workbook(source: str | Path, dest: str | Path, rows: list[dict]) -> int:
    """Copy the uploaded workbook and append calculated RWA after each row."""
    src = Path(source)
    book = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    written = 0
    try:
        columns: dict[str, int] = {}
        for item in rows:
            sheet_name = str(item.get("sheet") or "")
            if sheet_name not in book.sheetnames:
                continue
            ws = book[sheet_name]
            column = columns.get(sheet_name)
            if column is None:
                column = _next_value_column(ws)
                columns[sheet_name] = column
            header_row = int(item["header_row"])
            header = ws.cell(header_row, column)
            if header.value in (None, ""):
                header.value = "计算RWA"
                header.font = Font(bold=True)
            value = item.get("rwa")
            if not isinstance(value, (int, float)):
                continue
            cell = ws.cell(int(item["excel_row"]), column, float(value))
            cell.number_format = "#,##0.00"
            written += 1
        book.save(dest)
    finally:
        book.close()
    return written


def _next_value_column(ws) -> int:
    last = 0
    limit = min(ws.max_column or 1, 512)
    for row in ws.iter_rows(max_col=limit):
        for cell in reversed(row):
            if cell.value not in (None, ""):
                last = max(last, cell.column)
                break
    if last >= 512:
        raise ValueError("这张表没有空列可以写入计算 RWA")
    return last + 1


def _money(amount: float | None, currency: str | None) -> str:
    if amount is None or currency is None:
        return "—"
    unit = "美元" if currency == "USD" else "澳门元"
    return f"{amount:,.2f} {unit}"


def _col_index(letters: str) -> int:
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - 64)
    return index - 1
