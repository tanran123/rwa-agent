"""用大模型把底稿每一张明细的字段对上，再按权重法决策树计算风险加权资产。"""

from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, date
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from rwa_agent.exposure_class import class_fields, prototype_flow
from rwa_agent.mapper import _complete, _extract_json, suggest_complete_map
from rwa_agent.working_paper import PAPER_PATH, _num, _yi

CACHE_PATH = Path(__file__).resolve().parents[1] / "knowledge" / "底稿模型计算.json"
_PAPER_CONTEXT: dict[str, Any] = {}

SHEETS = [
    {"sheet": "term-Loan", "header_row": 14, "off": False},
    {"sheet": "Mortgage", "header_row": 5, "off": False},
    {"sheet": "BILLS", "header_row": 9, "off": False},
    {"sheet": "212 222 存放同业", "header_row": 13, "off": False},
    {"sheet": "Bonds", "header_row": 12, "off": False},
    {"sheet": "OD", "header_row": 11, "off": False},
    {"sheet": "額度詳情表", "header_row": 14, "off": True},
    {"sheet": "信用卡未使用額度表", "header_row": 8, "off": True},
    {"sheet": "TS-D-005", "header_row": 12, "off": True, "notional": True},
]


def load_or_compute(force: bool = False) -> dict[str, Any]:
    source = PAPER_PATH
    if not source.is_file():
        raise FileNotFoundError("找不到信用风险底稿：" + str(source))
    if not force and CACHE_PATH.is_file() and CACHE_PATH.stat().st_mtime >= source.stat().st_mtime:
        cached = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
        if cached.get("file") == source.name and cached.get("sheets") and cached.get("process_rows") is not None and cached.get("sheet_plan"):
            return cached
    result = compute_working_paper(source)
    CACHE_PATH.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return result


def write_result_workbook(data: dict[str, Any], dest) -> None:
    """把算完的敞口、风险加权资产和字段对应写成可下载的 Excel。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A5F4A")
    thin = Border(
        left=Side(style="thin", color="D9E4DE"),
        right=Side(style="thin", color="D9E4DE"),
        top=Side(style="thin", color="D9E4DE"),
        bottom=Side(style="thin", color="D9E4DE"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    def paint(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append(row)
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
            for cell in row:
                cell.border = thin
                cell.alignment = wrap
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    ead_yi = float(data.get("ead_yi") or 0)
    rwa_yi = float(data.get("rwa_yi") or 0)
    unpriced_yi = float(data.get("unpriced_yi") or 0)
    summary = book.active
    summary.title = "汇总"
    paint(summary, ["项目", "数值", "单位"], [
        ["底稿", data.get("file") or "", ""],
        ["数据日期", data.get("as_of") or "", ""],
        ["敞口", round(ead_yi, 4), "亿元"],
        ["敞口", round(ead_yi * 100_000_000, 2), "澳门元"],
        ["风险加权资产", round(rwa_yi, 4), "亿元"],
        ["风险加权资产", round(rwa_yi * 100_000_000, 2), "澳门元"],
        ["风险加权资产", round(float(data.get("rwa_thousand") or 0), 2), "千澳门元"],
        ["权重未定敞口", round(unpriced_yi, 4), "亿元"],
        ["已对应字段", data.get("mapped_fields") or 0, "个"],
        ["分类字段", data.get("field_count") or 0, "个"],
        ["说明", data.get("note") or "", ""],
    ], [22, 42, 14])

    result = book.create_sheet("分类结果")
    result_rows = []
    for line in data.get("lines") or []:
        ead = float(line.get("ead_yi") or 0)
        rwa = float(line.get("rwa_yi") or 0)
        weight = line.get("weight")
        result_rows.append([
            line.get("part") or "",
            line.get("section") or "",
            line.get("label") or "",
            line.get("weight_text") or ("未定权重" if weight is None else ""),
            line.get("model_reason") or "",
            line.get("count") or 0,
            round(ead, 4),
            round(rwa, 4),
            round(ead * 100_000_000, 2),
            round(rwa * 100_000_000, 2),
        ])
    paint(result, ["明细表", "风险暴露", "细分", "权重", "模型结论", "笔数", "敞口（亿元）", "风险加权资产（亿元）", "敞口（澳门元）", "风险加权资产（澳门元）"], result_rows, [24, 28, 42, 14, 36, 10, 16, 24, 20, 26])

    by_sheet = book.create_sheet("分表")
    sheet_rows = []
    for sheet in data.get("sheets") or []:
        sheet_rows.append([
            sheet.get("sheet") or "",
            sheet.get("rows") or 0,
            sheet.get("mapped") or 0,
            round(float(sheet.get("ead_yi") or 0), 4),
            round(float(sheet.get("rwa_yi") or 0), 4),
        ])
    paint(by_sheet, ["明细表", "笔数", "对上字段", "敞口（亿元）", "风险加权资产（亿元）"], sheet_rows, [28, 12, 14, 16, 24])

    fields = book.create_sheet("字段对应")
    field_rows = []
    for field in data.get("coverage") or []:
        hits = field.get("hits") or []
        field_rows.append([
            field.get("label") or "",
            field.get("id") or "",
            "\n".join(hits) if hits else "底稿没有",
        ])
    paint(fields, ["分类字段", "字段", "底稿对应"], field_rows, [36, 32, 72])
    book.save(dest)


def write_annotated_source(data: dict[str, Any], dest, source: Path | None = None) -> None:
    """在源工作簿每一行后面加上风险暴露类型、权重、敞口、风险加权资产和判定说明。"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    rows = data.get("process_rows") or []
    path = Path(source) if source else PAPER_PATH
    if not rows or not path.is_file():
        write_result_workbook(data, dest)
        return
    book = openpyxl.load_workbook(path)
    header_rows = {spec["sheet"]: spec["header_row"] for spec in SHEETS}
    titles = ["风险暴露类型", "权重", "敞口", "风险加权资产", "判定说明"]
    grouped: dict[str, list[dict]] = {}
    extra: list[dict] = []
    for row in rows:
        sheet = str(row.get("sheet") or "")
        excel_row = row.get("excel_row")
        if isinstance(excel_row, str) and excel_row.isdigit():
            excel_row = int(excel_row)
        if sheet in book.sheetnames and isinstance(excel_row, int):
            grouped.setdefault(sheet, []).append({**row, "excel_row": excel_row})
        else:
            extra.append(row)
    fill = PatternFill("solid", fgColor="1A5F4A")
    font = Font(bold=True, color="FFFFFF")
    for sheet, items in grouped.items():
        ws = book[sheet]
        header = next((int(item["header_row"]) for item in items if isinstance(item.get("header_row"), int)), 0)
        if header < 1:
            header = header_rows.get(sheet) or 1
        start = 0
        last = 1
        for cell in ws[header]:
            if cell.value == titles[0]:
                start = cell.column
            if cell.value not in (None, ""):
                last = cell.column
        if start < 1:
            start = last + 1
            for offset, title in enumerate(titles):
                cell = ws.cell(header, start + offset, title)
                cell.fill = fill
                cell.font = font
        for item in items:
            excel_row = int(item["excel_row"])
            ws.cell(excel_row, start, item.get("exposure") or "")
            ws.cell(excel_row, start + 1, item.get("weight_text") or "")
            ead = item.get("ead")
            rwa = item.get("rwa")
            if isinstance(ead, (int, float)):
                ws.cell(excel_row, start + 2, round(float(ead), 2))
            if isinstance(rwa, (int, float)):
                ws.cell(excel_row, start + 3, round(float(rwa), 2))
            ws.cell(excel_row, start + 4, item.get("trace") or "")
    if extra:
        supplement = book.create_sheet("计算补充")
        supplement.append(["来源", "名称", "风险暴露类型", "权重", "敞口", "风险加权资产", "判定说明"])
        for cell in supplement[1]:
            cell.font = font
            cell.fill = fill
        for item in extra:
            supplement.append([
                item.get("sheet") or "",
                item.get("name") or "",
                item.get("exposure") or "",
                item.get("weight_text") or "",
                item.get("ead"),
                item.get("rwa"),
                item.get("trace") or "",
            ])
    book.save(dest)


def compute_working_paper(path: Path | None = None) -> dict[str, Any]:
    import openpyxl

    source = Path(path) if path else PAPER_PATH
    plan_book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        specs, plan_note = _resolve_sheet_plan(plan_book)
    finally:
        plan_book.close()
    if not specs:
        raise ValueError(plan_note or "没有认出逐笔明细")
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    fields = _model_fields()
    prepared = []
    try:
        for spec in specs:
            headers, samples, rows = _read_sheet(book[spec["sheet"]], spec["header_row"])
            prepared.append({**spec, "headers": headers, "samples": samples, "rows": rows})
        prepared.extend(_supplement_sheets(book))
        global _PAPER_CONTEXT
        _PAPER_CONTEXT = _paper_context(book)
    finally:
        book.close()

    mapped: dict[str, dict] = {}
    errors: dict[str, str] = {}
    for item in prepared:
        binding = _strong_bind(item)
        if not _has_amount_column(binding) and item.get("rows"):
            try:
                binding = _map_sheet(item["sheet"], item["headers"], item["samples"], fields)
                binding["_by_model"] = True
            except Exception as exc:
                errors[item["sheet"]] = "模型没能判断这张表的列：" + str(exc)[:80]
        mapped[item["sheet"]] = binding

    categories = prototype_flow()["categories"]
    field_by_id = {item["id"]: item for item in class_fields()}
    ready = []
    for item in prepared:
        ready.append({**item, "mapping": mapped.get(item["sheet"]) or {}})
    bank_total = 0.0
    for item in ready:
        item["draft"] = _draft_rows(item, item["mapping"], field_by_id)
        bank_total += sum(amount for _facts, amount in item["draft"])
    _roll_up_debtors(ready)
    model_note = _fill_open_weights(ready, categories)
    sheets = []
    ead_sum = 0.0
    rwa_sum = 0.0
    unpriced = 0.0
    covered_ids: set[str] = set()
    lines: list[dict[str, Any]] = []
    for item in ready:
        sheet_result = _calculate_sheet(item, item["mapping"], categories, field_by_id, bank_total)
        if item["sheet"] in errors:
            sheet_result["error"] = errors[item["sheet"]]
        for field in sheet_result["fields"]:
            if field["status"] != "missing":
                covered_ids.add(field["id"])
        ead_sum += sheet_result["ead_mop"]
        rwa_sum += sheet_result["rwa_mop"]
        unpriced += sheet_result["unpriced_mop"]
        lines.extend(sheet_result["buckets"])
        sheets.append(sheet_result)
    coverage = []
    for fid, meta in field_by_id.items():
        hits = []
        for sheet in sheets:
            field = next((item for item in sheet["fields"] if item["id"] == fid), None)
            if not field or field["status"] == "missing":
                continue
            if field["status"] == "column":
                hits.append(sheet["sheet"] + " · " + field["column"] + " " + field["header"])
            elif field["status"] == "calc":
                hits.append(sheet["sheet"] + " · " + (field.get("header") or "按明细计算"))
            else:
                hits.append(sheet["sheet"] + " = " + str(field.get("const") or ""))
        coverage.append({"id": fid, "label": meta.get("label") or fid, "hits": hits})
    return {
        "file": source.name,
        "as_of": "20260930",
        "unit": "MOP",
        "field_count": len(fields),
        "mapped_fields": len(covered_ids),
        "ead_yi": _yi(ead_sum / 1000.0),
        "rwa_yi": _yi(rwa_sum / 1000.0),
        "unpriced_yi": _yi(unpriced / 1000.0),
        "rwa_thousand": rwa_sum / 1000.0,
        "note": (plan_note + " " if plan_note else "")
        + "现金、黄金走第五十七条，逾期和五级分类为次级、可疑、损失的走已违约，固定资产等走第八十一条的其他资产。存放中央银行按主权、未评级按 100%。敞口用账面余额，单一债务人按客户号加总。本地中小企只标企业划型，没有营业收入时不定 85%。外汇名义本金没有信用等值，不计入。"
        + ((" " + model_note) if model_note else ""),
        "sheet_plan": [{"sheet": item["sheet"], "header_row": item["header_row"], "reason": item.get("reason") or ""} for item in prepared],
        "model_note": model_note,
        "lines": lines,
        "process_rows": [row for sheet in sheets for row in (sheet.get("process_rows") or [])],
        "coverage": coverage,
        "sheets": [{key: value for key, value in sheet.items() if key not in {"fields", "process_rows"}} for sheet in sheets],
    }


def _model_fields() -> list[dict[str, Any]]:
    slim = []
    for field in class_fields():
        options = [str(opt["label"]) + "=" + str(opt["code"]) for opt in field.get("options") or []][:10]
        slim.append({
            "id": field["id"],
            "label": field["label"],
            "hint": (field.get("meaning") or "")[:70],
            "options": options,
        })
    return slim


def _has_amount_column(mapping: dict) -> bool:
    slot = mapping.get("single_exposure") or {}
    return slot.get("status") == "column"


def _amount_header_row(ws) -> int:
    words = ("balance", "余额", "餘額", "敞口", "equiv to", "额度", "額度")
    best = 0
    best_score = 0
    for index, row in enumerate(ws.iter_rows(max_row=30, max_col=40, values_only=True), 1):
        texts = []
        for value in row:
            if isinstance(value, str) and value.strip() and not value.startswith("="):
                texts.append(value.strip().lower())
        if len(texts) < 4:
            continue
        score = 0
        for text in texts:
            if "表" in text and not any(token in text.lower() for token in ("balance", "余额", "餘額")):
                continue
            lowered = text.lower()
            if any(word in lowered or word in text for word in words):
                score += 1
        if score > best_score:
            best = index
            best_score = score
    return best


def _fx_notional_header(ws, header_row: int) -> bool:
    texts = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, max_col=40, values_only=True):
        for value in row:
            if isinstance(value, str) and value.strip():
                texts.append(value.strip().lower())
        break
    blob = " ".join(texts)
    traded = any(token in blob for token in ("amount bought", "amount sold", "notional principal", "名义本金", "名義本金"))
    exposure = any(token in blob for token in ("value balance", "o/s bal", "信贷敞口", "信貸敞口", "剩余可使用", "剩餘可使用"))
    return traded and not exposure


def _header_width(ws, header_row: int) -> int:
    count = 0
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, max_col=40, values_only=True):
        for value in row:
            if isinstance(value, str) and value.strip() and not value.startswith("="):
                count += 1
        break
    return count


def _sheet_outline(ws) -> list[dict]:
    """挑出像表头的行给模型看，避免只看到标题就把整张明细漏掉。"""
    rows = []
    fallback = []
    for index, row in enumerate(ws.iter_rows(max_row=30, max_col=18, values_only=True), 1):
        cells = []
        for value in row:
            text = _sample(value)
            if text:
                cells.append(text[:28])
        if len(cells) >= 2 and len(fallback) < 2:
            fallback.append({"row": index, "cells": cells[:8]})
        if len(cells) >= 4:
            rows.append({"row": index, "cells": cells[:8]})
        if len(rows) >= 3:
            break
    return rows or fallback


def _plan_system() -> str:
    return (
        "用户上传了一本信用风险工作簿，用来计算 RWA。它不一定是某一家银行的固定底稿。"
        "请根据每张表的名称和前几行内容，判断哪些是要逐笔计算的明细，哪些是汇总、目录、评级对照或资产负债表。"
        "只输出一个 JSON 对象，不要解释。"
        "格式是 {\"note\":\"一句：这本表里哪些是明细\","
        "\"sheets\":[{\"sheet\":\"表名\",\"header_row\":表头行号,\"off\":false,\"notional\":false,\"skip\":false,\"reason\":\"一句\"}]}。"
        "header_row 从 1 开始，必须是真正的列名行，不是标题、单位或空行。"
        "只选带逐笔金额的明细：贷款、债券、存放、票据、透支、承诺额度和信用卡未用额度。"
        "客户名单、中小企名单、抵押品清单、评级对照、限额、科目对照、监管汇总和资产负债表不要逐笔算，skip 为 true。"
        "只有外汇名义本金、没有信用等值金额的表，notional 为 true 且 skip 为 true。"
        "表外承诺、未使用额度 off 为 true。"
        "不要编造工作簿里没有的表名。"
    )


def plan_workbook_sheets(book) -> tuple[list[dict], str]:
    """把整本工作簿的表名和表头样例交给模型，由它决定哪些表要逐笔计算。"""
    outlines = [{"sheet": name, "rows": _sheet_outline(book[name])} for name in book.sheetnames]
    raw = _complete(_plan_system(), json.dumps({"sheets": outlines}, ensure_ascii=False), timeout=120)
    parsed = _extract_json(raw)
    names = set(book.sheetnames)
    specs = []
    skipped = []
    seen: set[str] = set()
    for item in parsed.get("sheets") or []:
        if not isinstance(item, dict):
            continue
        sheet = str(item.get("sheet") or "")
        if sheet not in names or sheet in seen:
            continue
        seen.add(sheet)
        reason = str(item.get("reason") or "").strip()[:80]
        if item.get("skip") is True or item.get("notional") is True:
            if reason:
                skipped.append(sheet + "：" + reason)
            continue
        try:
            header_row = int(item.get("header_row"))
        except (TypeError, ValueError):
            continue
        if header_row < 1 or _header_width(book[sheet], header_row) < 4:
            continue
        detected = _amount_header_row(book[sheet])
        if not detected or _fx_notional_header(book[sheet], detected):
            continue
        off_name = any(token in sheet for token in ("额度", "額度", "未使用", "承诺", "承諾", "表外"))
        specs.append({
            "sheet": sheet,
            "header_row": detected,
            "off": off_name,
            "reason": reason,
        })
    note = str(parsed.get("note") or "").strip()[:180]
    if skipped:
        note = (note + " " if note else "") + "模型跳过 " + "；".join(skipped[:6])
    skipped_names = set()
    for item in parsed.get("sheets") or []:
        if isinstance(item, dict) and (item.get("skip") is True or item.get("notional") is True):
            skipped_names.add(str(item.get("sheet") or ""))
    specs = [item for item in specs if _amount_header_row(book[item["sheet"]])]
    chosen = {item["sheet"] for item in specs}
    added = []
    for name in book.sheetnames:
        if name in chosen or name in skipped_names:
            continue
        header = _amount_header_row(book[name])
        if not header or _fx_notional_header(book[name], header):
            continue
        specs.append({
            "sheet": name,
            "header_row": header,
            "off": any(token in name for token in ("额度", "額度", "未使用", "承诺", "承諾", "表外")),
            "reason": "表头里有金额列，补进逐笔明细",
        })
        added.append(name)
    if added:
        note = (note + " " if note else "") + "另按金额列补上 " + "、".join(added[:8])
    return specs, note


def _resolve_sheet_plan(book) -> tuple[list[dict], str]:
    known = [spec for spec in SHEETS if spec["sheet"] in book.sheetnames]
    try:
        planned, note = plan_workbook_sheets(book)
    except Exception as exc:
        planned, note = [], "模型暂时没能判断表结构，改用这本表上能认出的明细。" + str(exc)[:60]
    if planned:
        lead = "模型判断明细表：" + "、".join(item["sheet"] for item in planned[:8]) + "。"
        return planned, (lead + " " + note).strip()
    if known:
        return known, note or "模型没有给出可用的明细表，这本已知底稿仍按其原表明细计算。"
    return [], note or "模型没有从这本工作簿里认出逐笔明细。"


def _strong_bind(spec: dict) -> dict[str, dict]:
    """这本底稿的列名是固定的，能对上的字段直接绑列，对不上的留空。"""
    headers = spec["headers"]
    found: dict[str, dict] = {}
    name = spec["sheet"]

    def put(fid: str, slot: dict | None) -> None:
        if slot:
            found[fid] = slot

    def col(*needles: str, exact: bool = False) -> dict | None:
        for item in headers:
            header = item["header"]
            lowered = header.lower()
            for needle in needles:
                if exact and (header.strip() == needle or lowered.strip() == needle.lower()):
                    return {"status": "column", "column": item["column"], "reason": header}
                if not exact and (needle in header or needle.lower() in lowered):
                    return {"status": "column", "column": item["column"], "reason": header}
        return None

    def const(code: str, reason: str) -> dict:
        return {"status": "const", "const": code, "reason": reason}

    if not spec.get("notional"):
        put("single_exposure", col("Value Balance(MOP)", exact=True) or col("O/S BAL_MOP", exact=True) or col("Equiv to MOP", exact=True) or col("信貸敞口_MOP", "剩余可使用额度（澳门元）"))
    put("counterparty_name", col("Full Name", exact=True) or col("Customer Name", exact=True) or col("NAME", exact=True) or col("客戶名", "客户名", "中文名"))
    put("counterparty_type", col("Cust Type", exact=True) or col("INDUSTRY", exact=True))
    put("issuer_rating", col("RATING", exact=True))
    put("country_rating", col("COUNTRY Limit-short", exact=True))
    put("is_defaulted", col("五級分類", "五级分类") or col("DPD", exact=True) or col("逾期天数"))
    put("has_property_collateral", col("Description1", exact=True) or col("抵押品類型", "COLL. TYPE"))
    put("property_uses", col("Description1", exact=True) or col("抵押品類型"))
    put("provision_ratio", col("TOTAL ECL", exact=True))
    if _date_pair(headers):
        found["original_maturity_months"] = {"status": "calc", "header": "起息日到到期日", "reason": "到期日减起息日，换成月"}
    else:
        put("original_maturity_months", col("Original Period", "BY PERIOD"))
    found["instrument_type"] = const("claim", "这张明细是债权")
    found["treat_as_equity"] = const("no", "这张明细不是股权")
    found["is_subordinated"] = const("no", "这张明细不是次级债权")
    if name in _LOAN_SHEETS:
        found["fi_bank_applicable"] = const("not_applicable", "贷款、额度和信用卡不形成对商业银行的债权")
        found["is_credit_card"] = const("no", "不是信用卡")
    if "存放" in name:
        found["counterparty_type"] = const("commercial_bank", "212/222 存放同业")
        found["fi_bank_applicable"] = const("applicable", "存放同业适用商业银行")
        found["is_credit_card"] = const("no", "不是信用卡")
    if "信用卡" in name:
        found["is_credit_card"] = const("yes", "信用卡未使用额度")
        found["counterparty_type"] = const("individual", "持卡人是自然人")
        found["fi_bank_applicable"] = const("not_applicable", "信用卡不形成对商业银行的债权")
    if spec.get("notional"):
        found["single_exposure"] = {"status": "missing", "reason": "只有外汇名义本金，没有信用等值"}
    if name == "现金":
        found["instrument_type"] = const("cash", "资产负债表库存现金")
    elif name == "其他资产":
        found["instrument_type"] = const("other_asset", "没有交易对手的固定资产和递延费用")
    elif name == "存放中央银行":
        found["counterparty_type"] = const("foreign_central_bank", "资产负债表里的货币当局存款")
        found["issuer_rating"] = const("unrated", "底稿没有这家中央银行的外部评级")
    if "存放同业" in name:
        found["issuer_rating"] = {"status": "calc", "header": "同業rating 外部评级", "reason": "按客户名称对上 S&P、穆迪、惠誉"}
    if name in _LOAN_SHEETS:
        found["enterprise_size"] = {"status": "calc", "header": "本地中小企名单", "reason": "客户号在本地中小企名单里则为小型"}
        if found.get("is_defaulted", {}).get("status") != "column":
            found["is_defaulted"] = {"status": "calc", "header": "逾期科目或五级分类", "reason": "209 科目、户名含逾期，或五级分类为次级、可疑、损失"}
        if name in {"OD", "額度詳情表"} and found.get("counterparty_type", {}).get("status") != "column":
            found["counterparty_type"] = {"status": "calc", "header": "本地中小企或默认公司", "reason": "名单内或透支、额度没有客户类型时按公司"}
    return found


def _roll_up_debtors(ready: list[dict]) -> None:
    totals: dict[str, float] = {}
    for item in ready:
        for facts, amount in item["draft"]:
            key = str(facts.get("_debtor") or "")
            if key:
                totals[key] = totals.get(key, 0.0) + amount
    for item in ready:
        for facts, _amount in item["draft"]:
            key = str(facts.get("_debtor") or "")
            if key and key in totals:
                facts["single_exposure"] = totals[key]
                facts["_debtor_rolled"] = True


def _map_sheet(name: str, headers: list[dict], samples: list[dict], fields: list[dict]) -> dict[str, dict]:
    table = {"file": PAPER_PATH.name, "sheet": name, "rows": samples, "note": "只附了几行样例，列和整表一致。"}
    result = suggest_complete_map(headers, fields, table)
    return result["mapping"]


_NO_CONST = {
    "operating_normal", "no_major_violation_3y", "not_on_dishonest_list", "no_false_statements",
    "annual_audit_opinion", "no_default_3y", "loan_classification_normal", "guarantee_reasonable",
    "profit_tests_met", "no_other_material_risk", "transactor_clean", "meets_minimum_capital",
    "meets_other_capital", "bank_tier", "prudent", "manual_bank_grade", "revenue_mainly_central",
    "cet1_ratio", "leverage_ratio", "debt_ratio", "cashflow_dependent", "financial_industry",
    "state_subsidized_equity", "covered_pool_eligible", "covered_disclosure_ok",
}

_LOAN_SHEETS = {"term-Loan", "Mortgage", "BILLS", "OD", "額度詳情表", "信用卡未使用額度表"}

_AMOUNT_HEADERS = (
    "value balance(mop)",
    "o/s bal_mop",
    "equiv to mop",
    "信貸敞口_mop",
    "信贷敞口_mop",
    "剩余可使用额度（澳门元）",
    "剩餘可使用額度（澳門元）",
)


def _repair_mapping(spec: dict, mapping: dict[str, dict]) -> dict[str, dict]:
    """模型必须给每个字段一个结论；这里只改明显对错列的几项，避免抵押品被当成敞口。"""
    found = {fid: dict(slot) for fid, slot in mapping.items()}
    headers = spec["headers"]
    by_letter = {item["column"]: item["header"] for item in headers}
    for fid in list(found):
        slot = found[fid]
        if fid in _NO_CONST and slot.get("status") == "const":
            found[fid] = {"status": "missing", "reason": "底稿没有这一列，不能整表默认"}
    amount = _best_amount_column(headers)
    current = found.get("single_exposure") or {}
    current_header = by_letter.get(current.get("column") or "", "")
    if amount and not (current.get("status") == "column" and _good_amount_header(current_header)):
        found["single_exposure"] = {
            "status": "column",
            "column": amount[0],
            "reason": "账面余额 " + amount[1],
        }
    period = _period_column(headers)
    maturity = found.get("original_maturity_months") or {}
    maturity_header = by_letter.get(maturity.get("column") or "", "")
    if maturity.get("status") == "column" and not _period_header(maturity_header):
        if period:
            found["original_maturity_months"] = {
                "status": "column",
                "column": period[0],
                "reason": "期限天数 " + period[1],
            }
        else:
            found["original_maturity_months"] = {"status": "missing", "reason": "只有到期日，没有期限天数"}
    elif period and maturity.get("status") != "column":
        found["original_maturity_months"] = {
            "status": "column",
            "column": period[0],
            "reason": "期限天数 " + period[1],
        }
    name = spec["sheet"]
    if name in _LOAN_SHEETS:
        found["fi_bank_applicable"] = {"status": "const", "const": "not_applicable", "reason": "贷款、额度和信用卡不形成对商业银行的债权"}
    if "存放" in name:
        found["counterparty_type"] = {"status": "const", "const": "commercial_bank", "reason": "这张表是存放同业"}
        found["fi_bank_applicable"] = {"status": "const", "const": "applicable", "reason": "存放同业适用商业银行"}
        found["instrument_type"] = {"status": "const", "const": "claim", "reason": "存放同业是债权"}
    if "信用卡" in name:
        found["is_credit_card"] = {"status": "const", "const": "yes", "reason": "这张表是信用卡未使用额度"}
        found["instrument_type"] = {"status": "const", "const": "claim", "reason": "信用卡额度是债权"}
        found["counterparty_type"] = {"status": "const", "const": "individual", "reason": "信用卡持卡人按自然人"}
    if name == "額度詳情表":
        party = found.get("counterparty_type") or {}
        if party.get("status") == "const":
            found["counterparty_type"] = {"status": "missing", "reason": "额度表没有交易对手类型，不能整表写成主权"}
    if spec.get("notional"):
        found["single_exposure"] = {"status": "missing", "reason": "这张表只有外汇名义本金，没有信用等值"}
    return found


def _best_amount_column(headers: list[dict]) -> tuple[str, str] | None:
    ranked = []
    for item in headers:
        header = item["header"]
        if not _good_amount_header(header):
            continue
        lowered = header.lower()
        score = 0
        if "value balance(mop)" in lowered:
            score = 0
        elif "o/s bal_mop" in lowered:
            score = 1
        elif lowered.strip() == "equiv to mop":
            score = 2
        elif "信貸敞口_mop" in lowered or "信贷敞口_mop" in lowered:
            score = 3
        elif "剩余可使用额度" in header or "剩餘可使用額度" in header:
            score = 4
        else:
            score = 9
        ranked.append((score, item["column"], header))
    if not ranked:
        return None
    ranked.sort()
    return ranked[0][1], ranked[0][2]


def _good_amount_header(header: str) -> bool:
    if _bad_amount_header(header):
        return False
    lowered = header.lower()
    if lowered.strip() == "equiv to mop":
        return True
    return any(token in lowered or token in header for token in _AMOUNT_HEADERS)


def _bad_amount_header(header: str) -> bool:
    lowered = header.lower()
    tokens = ("collateral", "抵押", "rwa", "权重", "權重", "interest", "利息", "ccf", "ecl", "acquisition", "notional")
    return any(token in lowered or token in header for token in tokens)


def _period_column(headers: list[dict]) -> tuple[str, str] | None:
    for item in headers:
        if _period_header(item["header"]):
            return item["column"], item["header"]
    return None


def _period_header(header: str) -> bool:
    lowered = header.lower()
    if "maturity" in lowered or "到期" in header or "value date" in lowered:
        return False
    return "rest period" in lowered or "original period" in lowered or "by period" in lowered


def _header_column(headers: list[dict], names: tuple[str, ...]) -> tuple[str, str] | None:
    for item in headers:
        if any(name in item["header"] for name in names):
            return item["column"], item["header"]
    return None


def _read_sheet(ws, header_row: int) -> tuple[list[dict], list[dict], list[dict]]:
    header_cells: list[tuple[int, str]] = []
    for index, row in enumerate(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=80, values_only=True), 1):
        for col, value in enumerate(row, 1):
            if not isinstance(value, str):
                continue
            text = " ".join(value.replace("\n", " ").split())
            if text:
                header_cells.append((col, text[:80]))
        break
    letters = {col: get_column_letter(col) for col, _ in header_cells}
    headers = [{"column": letters[col], "header": text} for col, text in header_cells]
    rows: list[dict[str, Any]] = []
    samples: list[dict[str, str]] = []
    blank = 0
    for offset, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=header_row + 20000, max_col=80, values_only=True)):
        item = {"_excel_row": header_row + 1 + offset}
        for col, _text in header_cells:
            if col - 1 < len(row):
                item[letters[col]] = row[col - 1]
        if not any(value not in (None, "") for key, value in item.items() if not str(key).startswith("_")):
            blank += 1
            if blank >= 40 and rows:
                break
            continue
        blank = 0
        rows.append(item)
        if len(samples) < 2:
            samples.append({key: _sample(value) for key, value in item.items() if not str(key).startswith("_") and _sample(value)})
    for header in headers:
        picked = []
        for row in rows[:3]:
            text = _sample(row.get(header["column"]))
            if text and text not in picked:
                picked.append(text)
        header["samples"] = picked[:2]
    return headers, samples, rows


def _sample(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text or text.startswith("="):
        return ""
    return text[:36]


def _draft_rows(spec: dict, mapping: dict[str, dict], field_by_id: dict) -> list[tuple[dict, float]]:
    header_by_letter = {item["column"]: item["header"] for item in spec["headers"]}
    ccf_letter = next((item["column"] for item in spec["headers"] if item["header"].strip().upper() == "CCF"), "")
    draft = []
    if spec.get("notional"):
        return draft
    for row in spec["rows"]:
        facts = _facts_from_row(row, mapping, field_by_id, spec["sheet"], header_by_letter)
        amount = facts.pop("_amount", None)
        ccf = _num(row.get(ccf_letter)) if ccf_letter else None
        if amount is None or amount == 0:
            continue
        if not _row_has_identity(row, header_by_letter):
            continue
        if spec["off"] and ccf is not None and 0 < ccf <= 1:
            amount = amount * ccf
        elif spec["off"]:
            months = facts.get("original_maturity_months")
            if isinstance(months, (int, float)):
                amount = amount * (0.5 if months > 12 else 0.2)
            elif "信用卡" in spec["sheet"]:
                amount = amount * 0.5
        draft.append((facts, amount))
    return draft


def _calculate_sheet(spec: dict, mapping: dict[str, dict], categories: list, field_by_id: dict, bank_total: float) -> dict[str, Any]:
    header_by_letter = {item["column"]: item["header"] for item in spec["headers"]}
    fields_out = []
    for fid, meta in field_by_id.items():
        slot = mapping.get(fid) or {"status": "missing", "reason": "模型没有给出这一项"}
        fields_out.append({
            "id": fid,
            "label": meta.get("label") or fid,
            "status": slot.get("status") or "missing",
            "column": slot.get("column") or "",
            "header": header_by_letter.get(slot.get("column") or "", "") if slot.get("status") == "column" else (slot.get("header") or ""),
            "const": slot.get("const"),
            "reason": slot.get("reason") or "",
        })
    buckets: dict[str, dict[str, Any]] = {}
    ead_sum = 0.0
    rwa_sum = 0.0
    unpriced = 0.0
    process_rows: list[dict[str, Any]] = []
    draft = spec.get("draft") or []
    used_rows = len(draft)
    ltv_used = False
    for facts, amount in draft:
        if facts.get("_ltv_calc"):
            ltv_used = True
            facts.pop("_ltv_calc", None)
        debtor = facts.get("single_exposure")
        if isinstance(debtor, (int, float)) and bank_total:
            facts["exposure_share"] = float(debtor) / bank_total
        elif bank_total:
            facts["exposure_share"] = amount / bank_total
        hit = _classify(categories, facts)
        rw = hit["rw"]
        ead_sum += amount
        key = hit["subclass"] + "|" + (hit["rw_text"] or "未定")
        if facts.get("_model_reason"):
            key += "|model"
        bucket = buckets.get(key)
        if bucket is None:
            bucket = {
                "name": spec["sheet"] + " · " + hit["subclass"],
                "part": spec["sheet"],
                "section": hit["category"],
                "weight_text": hit["rw_text"] or "未定权重",
                "weight": rw,
                "ead_mop": 0.0,
                "rwa_mop": 0.0,
                "count": 0,
            }
            buckets[key] = bucket
        bucket["ead_mop"] += amount
        bucket["count"] += 1
        reason = str(facts.get("_model_reason") or "")
        if reason:
            reasons = bucket.setdefault("reasons", [])
            if reason not in reasons:
                reasons.append(reason)
        if isinstance(rw, float):
            bucket["rwa_mop"] += amount * rw
            rwa_sum += amount * rw
        else:
            unpriced += amount
        who = str(facts.get("counterparty_name") or "")
        if not who:
            debtor = str(facts.get("_debtor") or "")
            who = debtor[5:] if debtor.startswith("name:") else debtor
        exposure = hit["subclass"] or hit["category"] or "未归类"
        if hit["category"] and hit["category"] != exposure:
            exposure = hit["category"] + " · " + exposure
        process_rows.append({
            "sheet": spec["sheet"],
            "header_row": spec.get("header_row"),
            "excel_row": facts.get("_excel_row"),
            "name": who[:60],
            "exposure": exposure[:80],
            "weight": rw,
            "weight_text": hit["rw_text"] or "未定权重",
            "ead": amount,
            "rwa": amount * rw if isinstance(rw, float) else None,
            "trace": _trace_text(hit, facts),
        })
    lines = []
    for bucket in buckets.values():
        lines.append({
            "part": bucket["part"],
            "section": bucket["section"],
            "label": bucket["name"],
            "weight": bucket["weight"],
            "weight_text": bucket["weight_text"],
            "model_reason": "；".join(bucket.get("reasons") or []),
            "ead_yi": _yi(bucket["ead_mop"] / 1000.0),
            "rwa_yi": _yi(bucket["rwa_mop"] / 1000.0),
            "count": bucket["count"],
        })
    if any(facts.get("_debtor_rolled") for facts, _amount in draft):
        for field in fields_out:
            if field["id"] == "single_exposure" and field["status"] == "column":
                field["header"] = (field.get("header") or "账面余额") + "，按客户号加总"
    if ltv_used:
        for field in fields_out:
            if field["id"] == "ltv" and field["status"] == "missing":
                field["status"] = "calc"
                field["header"] = "账面余额÷抵押品价值"
                field["reason"] = "用账面余额除以抵押品价值"
    if bank_total:
        for field in fields_out:
            if field["id"] == "exposure_share" and field["status"] == "missing":
                field["status"] = "calc"
                field["header"] = "该笔敞口÷全部明细敞口"
                field["reason"] = "按全行明细合计计算占比"
    mapped = sum(1 for item in fields_out if item["status"] != "missing")
    return {
        "sheet": spec["sheet"],
        "rows": used_rows,
        "mapped": mapped,
        "field_count": len(fields_out),
        "fields": fields_out,
        "ead_mop": ead_sum,
        "rwa_mop": rwa_sum,
        "unpriced_mop": unpriced,
        "ead_yi": _yi(ead_sum / 1000.0),
        "rwa_yi": _yi(rwa_sum / 1000.0),
        "buckets": lines,
        "process_rows": process_rows,
    }


def _supplement_sheets(book) -> list[dict]:
    """资产负债表上没有交易明细的现金、货币当局存款和其他资产。"""
    if "FS-ASSETS" not in book.sheetnames:
        return []
    assets = book["FS-ASSETS"]
    cash = _balance_line(assets, "現金")
    central = _balance_line(assets, "AMCM")
    others = [
        ("不動產", _balance_line(assets, "不動產")),
        ("設備", _balance_line(assets, "設備")),
        ("遞延費用", _balance_line(assets, "遞延費用")),
        ("開辦費用", _balance_line(assets, "開辦費用")),
        ("未完成不動產", _balance_line(assets, "未完成不動產")),
    ]
    sheets = []
    if cash:
        sheets.append(_synthetic("现金", [("库存现金", cash)]))
    if central:
        sheets.append(_synthetic("存放中央银行", [("AUTORIDADE MONETARIA DE MACAU", central)]))
    other_rows = [(label, amount) for label, amount in others if amount]
    if other_rows:
        sheets.append(_synthetic("其他资产", other_rows))
    return sheets


def _synthetic(name: str, rows: list[tuple[str, float]]) -> dict:
    headers = [
        {"column": "A", "header": "Customer Name"},
        {"column": "B", "header": "Value Balance(MOP)"},
    ]
    data = [{"A": label, "B": amount} for label, amount in rows]
    return {"sheet": name, "header_row": 1, "off": False, "headers": headers, "samples": [], "rows": data}


def _balance_line(ws, prefix: str) -> float:
    for row in ws.iter_rows(min_row=6, max_row=24, max_col=4, values_only=True):
        label = str(row[0] or "").strip()
        if label.startswith(prefix):
            return _num(row[3]) or 0.0
    return 0.0


def _paper_context(book) -> dict[str, Any]:
    sme: set[str] = set()
    if "本地中小企" in book.sheetnames:
        ws = book["本地中小企"]
        for row in ws.iter_rows(min_row=2, max_row=80, max_col=4, values_only=True):
            if str(row[3] or "").strip().upper() == "Y" and row[0] not in (None, ""):
                sme.add(str(row[0]).strip())
    principal = 0.0
    provision = 0.0
    if "past due loan" in book.sheetnames:
        past = book["past due loan"]
        for row in past.iter_rows(min_row=1, max_row=20, max_col=5, values_only=True):
            label = str(row[2] or "")
            if "權重" not in label and "权重" not in label:
                continue
            base = _num(row[3]) or 0.0
            ecl = _num(row[4]) or 0.0
            if base > 0:
                principal += base
                provision += ecl
    ratings: dict[str, str] = {}
    if "同業rating" in book.sheetnames:
        peers = book["同業rating"]
        for row in peers.iter_rows(min_row=2, max_row=80, max_col=5, values_only=True):
            name = str(row[1] or "").strip()
            if not name:
                continue
            found = []
            for cell in row[2:5]:
                code = _rating_code(str(cell or ""))
                if code and code != "unrated":
                    found.append(code)
            if found:
                ratings[_party_key(name)] = max(found, key=_RATING_RANK.index)
    ratio = (provision / principal) if principal else None
    return {"sme": sme, "default_provision": ratio, "ratings": ratings}


def _party_key(name: str) -> str:
    return " ".join(name.upper().replace(".", " ").split())


def _header_value(row: dict, header_by_letter: dict[str, str], *names: str):
    wanted = {name.upper() for name in names}
    for letter, header in header_by_letter.items():
        if header.strip().upper() in wanted:
            return row.get(letter)
    return None


def _apply_paper_context(facts: dict, row: dict, header_by_letter: dict[str, str], sheet: str, amount: float | None) -> None:
    context = _PAPER_CONTEXT or {}
    account = str(_header_value(row, header_by_letter, "Account Name") or "")
    gl = str(_header_value(row, header_by_letter, "GL A/C No.") or "")
    product = str(_header_value(row, header_by_letter, "PRODUCT TYPE", "Product Code") or "")
    upper = (account + " " + product).upper()
    dpd = _num(_header_value(row, header_by_letter, "DPD"))
    if dpd is not None and dpd > 90:
        facts["is_defaulted"] = "yes"
    loan_default = gl.startswith("209") or "OVERDUE" in upper or "逾期" in account
    if loan_default:
        facts["is_defaulted"] = "yes"
        if "provision_ratio" not in facts:
            ratio = context.get("default_provision")
            if isinstance(ratio, float):
                facts["provision_ratio"] = ratio
    if "CREDIT CARD" in upper or gl.startswith("207"):
        facts["is_credit_card"] = "yes"
        if facts.get("counterparty_type") not in {"individual", "corporate"}:
            facts["counterparty_type"] = "individual"
    debtor = str(facts.get("_debtor") or "")
    if debtor in context.get("sme", ()):
        facts["enterprise_size"] = "small"
        if "counterparty_type" not in facts:
            facts["counterparty_type"] = "corporate"
    local = str(_header_value(row, header_by_letter, "本地中小企") or "").strip().upper()
    if local == "Y":
        facts["enterprise_size"] = "small"
        if "counterparty_type" not in facts:
            facts["counterparty_type"] = "corporate"
    personal = _header_value(row, header_by_letter, "濠人貸")
    if sheet in {"OD", "額度詳情表"} and "counterparty_type" not in facts:
        facts["counterparty_type"] = "individual" if personal in {True, "Y", "TRUE", "是"} else "corporate"
    if "存放" in sheet and "issuer_rating" not in facts:
        party = _party_key(str(facts.get("counterparty_name") or ""))
        rated = (context.get("ratings") or {}).get(party)
        if rated:
            facts["issuer_rating"] = rated
    if isinstance(amount, float) and facts.get("provision_ratio") and facts["provision_ratio"] > 1:
        facts["provision_ratio"] = facts["provision_ratio"] / amount


def _facts_from_row(row: dict, mapping: dict[str, dict], field_by_id: dict, sheet: str, header_by_letter: dict[str, str]) -> dict[str, Any]:
    facts: dict[str, Any] = {}
    amount = None
    for fid, meta in field_by_id.items():
        slot = mapping.get(fid) or {}
        status = slot.get("status")
        raw = None
        header = ""
        if status == "column":
            raw = row.get(slot.get("column") or "")
            header = header_by_letter.get(slot.get("column") or "", "")
        elif status == "const":
            raw = slot.get("const")
        if raw in (None, "") :
            if status == "column" and fid in {"issuer_rating", "bond_rating", "country_rating"}:
                facts[fid] = "unrated"
            continue
        value = _coerce(fid, raw, meta.get("options") or [], header)
        if value is None:
            continue
        if fid == "single_exposure" and isinstance(value, float):
            amount = value
        facts[fid] = value
    ratio = facts.get("provision_ratio")
    if isinstance(ratio, float) and ratio > 1 and amount:
        facts["provision_ratio"] = ratio / amount
    if "is_defaulted" not in facts:
        dpd = _find_number(row, mapping, "is_defaulted")
        if dpd is not None and dpd > 90:
            facts["is_defaulted"] = "yes"
    uses = _property_uses(row, header_by_letter)
    if uses:
        facts["has_property_collateral"] = "yes"
        facts["property_uses"] = uses
        collateral = []
        for group in _collateral_groups(row, header_by_letter):
            desc = str(group.get("desc") or "").strip()
            if desc and _looks_property(desc):
                collateral.append(desc[:60])
        if collateral:
            facts["_collateral"] = " | ".join(collateral)[:160]
        text = _text_blob(row).upper()
        facts["cashflow_dependent"] = "yes" if any(token in text for token in ("出租", "租金", "收租", "RENTAL", "INVESTMENT")) else "no"
    elif "has_property_collateral" in facts and facts["has_property_collateral"] != "yes":
        facts.pop("has_property_collateral", None)
        facts.pop("property_uses", None)
    months = _original_months(row, header_by_letter)
    if months:
        facts["original_maturity_months"] = months
    debtor = _debtor_of(row, header_by_letter)
    if debtor:
        facts["_debtor"] = debtor
    if "存放" in sheet:
        party = str(facts.get("counterparty_name") or "").upper()
        if any(token in party for token in ("CENTRAL BANK", "AMCM", "AUTORIDADE", "央行")):
            facts["counterparty_type"] = "foreign_central_bank"
    collateral = _property_collateral_mop(row, header_by_letter)
    if isinstance(amount, float) and collateral > 0:
        facts["ltv"] = amount / collateral
        facts["_ltv_calc"] = True
    if facts.get("counterparty_type") == "mdb":
        official = _mdb_official_name(_text_blob(row))
        if official:
            facts["counterparty_name"] = official
    if facts.get("counterparty_type") in {"foreign_sovereign", "cn_central_government"}:
        country = _country_text(row, header_by_letter)
        if _is_prc(country):
            facts["counterparty_type"] = "cn_central_government"
        elif country:
            facts["counterparty_type"] = "foreign_sovereign"
    _apply_paper_context(facts, row, header_by_letter, sheet, amount if isinstance(amount, float) else None)
    facts["_amount"] = amount if amount is not None else facts.get("single_exposure") if isinstance(facts.get("single_exposure"), float) else None
    facts["_excel_row"] = row.get("_excel_row")
    return facts


def _find_number(row: dict, mapping: dict, fid: str) -> float | None:
    slot = mapping.get(fid) or {}
    if slot.get("status") != "column":
        return None
    return _num(row.get(slot.get("column") or ""))


def _text_blob(row: dict) -> str:
    parts = []
    for value in row.values():
        if isinstance(value, str):
            parts.append(value)
    return " ".join(parts)


def _row_has_identity(row: dict, header_by_letter: dict[str, str]) -> bool:
    for letter, header in header_by_letter.items():
        lowered = header.lower()
        if not any(token in lowered or token in header for token in ("customer", "客户", "客戶", "name", "gl a/c", "contract", "isin", "卡号", "证件", "ci")):
            continue
        if "country" in lowered or "industry" in lowered:
            continue
        value = row.get(letter)
        if isinstance(value, str) and value.strip():
            return True
        if isinstance(value, (int, float)) and not isinstance(value, bool) and value != 0:
            return True
    return False


_MDB_OFFICIAL = (
    ("IBRD", "世界银行集团"),
    ("IFC", "世界银行集团"),
    ("ASIAN DEVELOPMENT", "亚洲开发银行"),
    ("ASIA ", "亚洲开发银行"),
    ("EBRD", "欧洲复兴开发银行"),
    ("IADB", "泛美开发银行"),
    ("EIB", "欧洲投资银行"),
    ("NIB", "北欧投资银行"),
    ("ISDB", "伊斯兰开发银行"),
    ("COE", "欧洲理事会开发银行"),
)


def _mdb_official_name(text: str) -> str:
    upper = text.upper()
    for token, official in _MDB_OFFICIAL:
        if token in upper:
            return official
    return ""


def _looks_residential(text: str) -> bool:
    return any(token in text.upper() for token in ("RESIDENTIAL", "MORTGAGE", "住宅", "按揭", "物業", "物业"))


def _looks_property(text: str) -> bool:
    upper = text.upper()
    if any(token in upper for token in ("GUARANTEE", "LIVRANCA", "DEPOSIT", "保证", "擔保", "存单", "存單")):
        return False
    if re.search(r"\bTD\b", upper):
        return False
    return any(token in upper for token in ("RESIDENTIAL", "COMMERCIAL PROPERTY", "PROPERTY", "住宅", "商用", "物業", "物业", "按揭"))


def _collateral_groups(row: dict, header_by_letter: dict[str, str]) -> list[dict]:
    groups: dict[str, dict] = {}
    for letter, header in header_by_letter.items():
        desc = re.search(r"Description\s*(\d+)", header, re.I)
        if desc:
            groups.setdefault(desc.group(1), {})["desc"] = str(row.get(letter) or "")
        value = re.search(r"Collateral Value.*(\d+)\s*$", header, re.I)
        if value and "mop" in header.lower():
            groups.setdefault(value.group(1), {})["value"] = _num(row.get(letter))
    return list(groups.values())


def _property_uses(row: dict, header_by_letter: dict[str, str]) -> list[str]:
    uses = []
    for group in _collateral_groups(row, header_by_letter):
        text = group.get("desc") or ""
        if not _looks_property(text):
            continue
        if _looks_residential(text) and "residential" not in uses:
            uses.append("residential")
        elif "commercial" not in uses:
            uses.append("commercial")
    return uses


def _property_collateral_mop(row: dict, header_by_letter: dict[str, str]) -> float:
    total = 0.0
    for group in _collateral_groups(row, header_by_letter):
        if not _looks_property(group.get("desc") or ""):
            continue
        number = group.get("value")
        if isinstance(number, float) and number > 0:
            total += number
    return total


def _debtor_of(row: dict, header_by_letter: dict[str, str]) -> str:
    by_header = {header.strip(): letter for letter, header in header_by_letter.items()}
    for name in ("Customer", "CI", "客户号", "客戶號", "证件号码"):
        letter = by_header.get(name)
        if not letter:
            continue
        value = row.get(letter)
        if value not in (None, ""):
            return str(value).strip()
    for name in ("Full Name", "Customer Name", "NAME", "客戶名", "中文名"):
        letter = by_header.get(name)
        if not letter:
            continue
        value = row.get(letter)
        if isinstance(value, str) and value.strip():
            return "name:" + value.strip()
    return ""


def _pick_date_column(header_by_letter: dict[str, str], kind: str) -> str:
    exact = ""
    fuzzy = ""
    for letter, header in header_by_letter.items():
        lowered = header.strip().lower()
        if kind == "value":
            if lowered == "value date":
                exact = letter
            elif not fuzzy and ("value date" in lowered or header.strip() in {"起息日", "额度生效日", "額度生效日", "发卡日期", "發卡日期"}):
                fuzzy = letter
        elif lowered == "maturity date":
            exact = letter
        elif not fuzzy and "maturity" in lowered and "next" not in lowered:
            fuzzy = letter
        elif not fuzzy and header.strip() in {"到期日", "额度到期日", "額度到期日", "信用卡到期日"}:
            fuzzy = letter
    return exact or fuzzy


def _date_pair(headers: list[dict]) -> bool:
    letters = {item["column"]: item["header"] for item in headers}
    return bool(_pick_date_column(letters, "value") and _pick_date_column(letters, "maturity"))


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _original_months(row: dict, header_by_letter: dict[str, str]) -> float | None:
    start = _as_date(row.get(_pick_date_column(header_by_letter, "value")))
    end = _as_date(row.get(_pick_date_column(header_by_letter, "maturity")))
    if not start or not end or end <= start:
        return None
    return (end - start).days / 30.0


def _country_text(row: dict, header_by_letter: dict[str, str]) -> str:
    for letter, header in header_by_letter.items():
        upper = header.upper()
        if "COUNTRY" not in upper and "國別" not in header and "国别" not in header:
            continue
        if "CODE" in upper:
            continue
        value = row.get(letter)
        if isinstance(value, str) and value.strip() and not _num(value):
            return value.strip()
    return ""


def _is_prc(country: str) -> bool:
    upper = country.upper().strip()
    if upper in {"CN", "CHN", "PRC", "CHINA"} or upper.startswith("CN ") or upper.startswith("CHINA"):
        return True
    return any(token in country for token in ("中国", "中國"))


_SP_RATING = re.compile(r"(?:標普|标普|S&P|SP)\s*([A-D]{1,3}\s*[+-]?)")
_FITCH_RATING = re.compile(r"(?:惠譽|惠誉|Fitch)\s*([A-D]{1,3}\s*[+-]?)")
_MOODY = (
    ("Aaa", "AAA"), ("Aa1", "AA+"), ("Aa2", "AA"), ("Aa3", "AA-"),
    ("A1", "A+"), ("A2", "A"), ("A3", "A-"),
    ("Baa1", "BBB+"), ("Baa2", "BBB"), ("Baa3", "BBB-"),
    ("Ba1", "BB+"), ("Ba2", "BB"), ("Ba3", "BB-"),
    ("B1", "B+"), ("B2", "B"), ("B3", "B-"),
    ("Caa", "CCC"), ("Ca", "CC"),
)


_RATING_RANK = [
    "AAA", "AA+", "AA", "AA-", "A+", "A", "A-",
    "BBB+", "BBB", "BBB-", "BB+", "BB", "BB-", "B+", "B", "B-",
    "CCC", "CC", "C", "D",
]
_MOODY_MARK = re.compile(r"(?:穆迪|Moody's|Moodys)\s*([A-C][a-z0-9]+)")


def _rating_code(text: str) -> str | None:
    cleaned = " ".join(text.replace("／", "/").split())
    found: list[str] = []
    for pattern in (_SP_RATING, _FITCH_RATING):
        match = pattern.search(cleaned)
        if match:
            found.append(match.group(1).replace(" ", ""))
    moody = _MOODY_MARK.search(cleaned)
    if moody:
        token = moody.group(1)
        for raw, sp in _MOODY:
            if token.startswith(raw):
                found.append(sp)
                break
    if not found and re.fullmatch(r"[A-D]{1,3}[+-]?", cleaned):
        found.append(cleaned)
    if not found:
        for raw, sp in _MOODY:
            if raw in cleaned:
                found.append(sp)
                break
    ranked = [item for item in found if item in _RATING_RANK]
    if ranked:
        return max(ranked, key=_RATING_RANK.index)
    if cleaned.upper() in {"NR", "N/A", "NA", "UNRATED", "无", "無"}:
        return "unrated"
    return None


def _coerce(fid: str, raw: Any, options: list[dict], header: str) -> Any:
    if isinstance(raw, datetime):
        return None
    if isinstance(raw, date):
        return None
    if fid in {"single_exposure", "ltv", "provision_ratio", "original_maturity_months", "cet1_ratio", "leverage_ratio", "debt_ratio", "covered_oc_ratio", "revenue_last_year"}:
        number = _num(raw)
        if number is None:
            return None
        if fid == "ltv" and number > 3:
            return number / 100.0
        if fid == "provision_ratio" and number > 1:
            return number
        if fid == "original_maturity_months" and _period_header(header):
            return number / 30.0
        if fid == "original_maturity_months" and number > 600 and any(token in header.lower() for token in ("day", "period", "rest")):
            return number / 30.0
        return number
    text = str(raw).strip()
    if fid in {"issuer_rating", "bond_rating", "country_rating"}:
        rated = _rating_code(text)
        if rated:
            return rated
    for opt in options:
        if text == str(opt.get("code")) or text == str(opt.get("label")):
            code = opt.get("code")
            if fid == "property_uses":
                return [code]
            return code
    upper = text.upper()
    if fid == "counterparty_type":
        if upper in {"P", "INDIVIDUAL"} or "个人" in text or "自然人" in text:
            return "individual"
        if upper in {"BANK"} or "银行" in text or "同业" in text:
            return "commercial_bank"
        if upper in {"SOVEREIGN"} or "主权" in text or "财政部" in text:
            return "cn_central_government" if any(token in upper for token in ("CN", "MO", "中国", "澳门", "我國")) else "foreign_sovereign"
        if upper in {"MDBS", "MDB"}:
            return "mdb"
        if upper in {"PSES", "PSE", "MUNICIPAL", "GOV"}:
            return "pse"
        if upper in {"C", "CORP", "CORPORATE"} or "公司" in text or "企业" in text:
            return "corporate"
    if fid == "is_defaulted":
        if any(token in upper for token in ("LOSS", "SUBSTANDARD", "DOUBTFUL", "次级", "可疑", "损失", "违约", "TRUE", "逾期")):
            return "yes"
        if upper in {"PASS", "NORMAL", "FALSE"} or text in {"正常", "关注", "否", "N", "NO"}:
            return "no"
    if fid in {"has_property_collateral", "cashflow_dependent", "is_credit_card", "treat_as_equity", "is_subordinated"}:
        if text in {"Y", "YES", "是", "TRUE"}:
            return "yes"
        if text in {"N", "NO", "否", "FALSE"}:
            return "no"
    if fid == "property_uses":
        if _looks_residential(text):
            return ["residential"]
        if any(token in text for token in ("商用", "商业", "COMMERCIAL")):
            return ["commercial"]
    if fid == "instrument_type":
        if text in {"cash", "gold", "other_asset", "equity", "covered_bond", "claim"}:
            return text
        if "现金" in text:
            return "cash"
        if "黄金" in text:
            return "gold"
        if "股权" in text:
            return "equity"
        if "担保债券" in text or "COVERED" in upper:
            return "covered_bond"
        if "其他资产" in text or "固定资产" in text:
            return "other_asset"
        return "claim"
    return None


_OPEN_OPTIONS = {
    "bank_grade": ["a_plus", "a", "b", "c"],
    "property_side": ["residential", "commercial"],
    "default_provision": ["high", "low"],
}


def _fill_open_weights(ready: list[dict], categories: list) -> str:
    """规则树停住、权重还没数的敞口，交给模型在允许的选项里定，再回到原规则重算。"""
    groups: dict[str, dict] = {}
    bound: list[tuple[dict, str]] = []
    for item in ready:
        for facts, amount in item["draft"]:
            hit = _classify(categories, facts)
            if isinstance(hit.get("rw"), float):
                continue
            kind = _open_kind(hit, facts)
            if not kind:
                continue
            case_id = _open_id(kind, facts, item["sheet"])
            group = groups.get(case_id)
            if group is None:
                group = {
                    "id": case_id,
                    "kind": kind,
                    "options": _OPEN_OPTIONS[kind],
                    "stuck": hit.get("subclass") or hit.get("category") or "",
                    "ead_mop": 0.0,
                    "count": 0,
                    "names": [],
                    "ratings": [],
                    "months": [],
                    "collateral": [],
                    "sheets": [],
                }
                groups[case_id] = group
            group["ead_mop"] += float(amount or 0)
            group["count"] += 1
            _remember_open(group, facts, item["sheet"])
            bound.append((facts, case_id))
    if not groups:
        return ""
    try:
        choices = _ask_open_weights(list(groups.values()))
    except Exception as exc:
        return "规则没有定出权重的敞口没能交给模型：" + str(exc)[:160]
    priced = {"bank_grade": 0, "property_side": 0, "default_provision": 0}
    missed = 0
    for facts, case_id in bound:
        picked = choices.get(case_id) or {}
        choice = str(picked.get("choice") or "")
        kind = groups[case_id]["kind"]
        if not _apply_open_choice(facts, kind, choice):
            missed += 1
            continue
        hit = _classify(categories, facts)
        reason = str(picked.get("reason") or "").strip()
        if isinstance(hit.get("rw"), float):
            facts["_model_reason"] = reason or _choice_label(kind, choice)
            priced[kind] += 1
        else:
            missed += 1
    parts = []
    if priced["bank_grade"]:
        parts.append("商业银行 " + str(priced["bank_grade"]) + " 笔由模型选定资本档次")
    if priced["property_side"]:
        parts.append("居住和商用共同抵押 " + str(priced["property_side"]) + " 笔由模型选定一边")
    if priced["default_provision"]:
        parts.append("已违约但没有减值准备的 " + str(priced["default_provision"]) + " 笔由模型选定 100% 或 150%")
    if missed:
        parts.append(str(missed) + " 笔模型没有给出可用结论，仍未定权重")
    if not parts:
        return "规则没有定出权重时已发给模型，模型没有返回可用结论。"
    return "规则定不了权重的已发给模型，在规则允许的选项里选定后重算：" + "；".join(parts) + "。"


def _open_kind(hit: dict, facts: dict) -> str:
    text = str(hit.get("subclass") or "") + str(hit.get("category") or "")
    if "共同抵押" in text:
        return "property_side"
    if "已违约" in text and "provision_ratio" not in facts:
        return "default_provision"
    if "商业银行" in text or facts.get("counterparty_type") == "commercial_bank":
        return "bank_grade"
    return ""


def _open_id(kind: str, facts: dict, sheet: str) -> str:
    if kind == "default_provision":
        return "default_provision"
    if kind == "property_side":
        collateral = str(facts.get("_collateral") or facts.get("counterparty_name") or facts.get("_debtor") or sheet)
        return "property:" + collateral[:80]
    name = _party_key(str(facts.get("counterparty_name") or ""))
    if not name:
        debtor = str(facts.get("_debtor") or "")
        name = debtor[5:] if debtor.startswith("name:") else debtor
        name = _party_key(name)
    return "bank:" + (name or sheet)


def _remember_open(group: dict, facts: dict, sheet: str) -> None:
    if sheet not in group["sheets"]:
        group["sheets"].append(sheet)
    name = str(facts.get("counterparty_name") or "")
    if not name:
        debtor = str(facts.get("_debtor") or "")
        name = debtor[5:] if debtor.startswith("name:") else ""
    if name and name not in group["names"] and len(group["names"]) < 4:
        group["names"].append(name[:60])
    rating = str(facts.get("issuer_rating") or "")
    if rating and rating not in group["ratings"]:
        group["ratings"].append(rating)
    months = facts.get("original_maturity_months")
    if isinstance(months, (int, float)):
        group["months"].append(float(months))
    collateral = str(facts.get("_collateral") or "")
    if collateral and collateral not in group["collateral"]:
        group["collateral"].append(collateral)


def _ask_open_weights(groups: list[dict]) -> dict[str, dict]:
    payload = []
    for group in groups:
        months = group["months"]
        payload.append({
            "id": group["id"],
            "kind": group["kind"],
            "options": group["options"],
            "stuck": group["stuck"],
            "count": group["count"],
            "ead_yi": round(group["ead_mop"] / 100000000, 4),
            "sheets": group["sheets"],
            "names": group["names"],
            "external_rating": group["ratings"],
            "maturity_months": [round(min(months), 1), round(max(months), 1)] if months else [],
            "collateral": group["collateral"][:3],
        })
    raw = _complete(_open_weight_system(), json.dumps({"cases": payload}, ensure_ascii=False), timeout=120)
    parsed = _extract_json(raw)
    incoming = parsed.get("choices")
    if isinstance(incoming, dict):
        incoming = [{"id": key, **value} if isinstance(value, dict) else {"id": key, "choice": value} for key, value in incoming.items()]
    if not isinstance(incoming, list):
        incoming = []
    allowed = {group["id"]: set(group["options"]) for group in groups}
    found: dict[str, dict] = {}
    for item in incoming:
        if not isinstance(item, dict):
            continue
        case_id = str(item.get("id") or "")
        choice = str(item.get("choice") or "").strip()
        if case_id not in allowed or choice not in allowed[case_id]:
            continue
        found[case_id] = {"choice": choice, "reason": str(item.get("reason") or "").strip()[:80]}
    return found


def _open_weight_system() -> str:
    return (
        "有些信用风险敞口走到了权重法规则，但规则因为缺条件停住了，权重还是空的。"
        "你要替这些敞口把缺的条件定下来。只输出一个 JSON 对象，不要解释。"
        "格式是 {\"choices\":[{\"id\":\"原样照抄\",\"choice\":\"选项之一\",\"reason\":\"一句为什么这样定\"}]}。"
        "每一个 case 的 id 都必须出现，choice 只能是该 case 的 options 里的值，不能写权重百分比。"
        "bank_grade：底稿没有资本充足率。a_plus 表示核心一级资本充足率不低于 14% 且杠杆率不低于 5%，并满足最低资本和其他各级资本；"
        "a 表示满足最低资本和其他各级资本，但达不到 A+；b 表示只满足最低资本、不满足其他各级资本；c 表示不满足最低资本。"
        "外部评级只是参考，不能把 AA、A 直接当成资本档次。期限由规则自己处理，不要用档次去表示期限。"
        "property_side：居住和商用同时抵押，规则要求先选定一边。residential 归居住用，commercial 归商用。按抵押品描述选更贴近的一边。"
        "default_provision：已违约，但没有这笔的减值准备，规则分不出 100% 和 150%。"
        "high 表示按损失准备不低于账面 20%，权重 100%；low 表示按损失准备低于账面 20%，权重 150%。"
        "信用卡未使用额度通常没有单独计提，更接近准备不足。"
    )


def _apply_open_choice(facts: dict, kind: str, choice: str) -> bool:
    if kind == "bank_grade":
        if choice == "a_plus":
            facts["meets_minimum_capital"] = "yes"
            facts["meets_other_capital"] = "yes"
            facts["cet1_ratio"] = 0.14
            facts["leverage_ratio"] = 0.05
        elif choice == "a":
            facts["meets_minimum_capital"] = "yes"
            facts["meets_other_capital"] = "yes"
        elif choice == "b":
            facts["meets_minimum_capital"] = "yes"
            facts["meets_other_capital"] = "no"
        elif choice == "c":
            facts["meets_minimum_capital"] = "no"
        else:
            return False
        return True
    if kind == "property_side":
        if choice not in {"residential", "commercial"}:
            return False
        facts["has_property_collateral"] = "yes"
        facts["property_uses"] = [choice]
        return True
    if kind == "default_provision":
        if choice == "high":
            facts["provision_ratio"] = 0.2
        elif choice == "low":
            facts["provision_ratio"] = 0.0
        else:
            return False
        return True
    return False


def _choice_label(kind: str, choice: str) -> str:
    labels = {
        ("bank_grade", "a_plus"): "模型定为 A+ 级",
        ("bank_grade", "a"): "模型定为 A 级",
        ("bank_grade", "b"): "模型定为 B 级",
        ("bank_grade", "c"): "模型定为 C 级",
        ("property_side", "residential"): "模型选定居住用",
        ("property_side", "commercial"): "模型选定商用",
        ("default_provision", "high"): "模型按准备不低于 20%，权重 100%",
        ("default_provision", "low"): "模型按准备低于 20%，权重 150%",
    }
    return labels.get((kind, choice), "模型选定")


def _classify(categories: list, facts: dict[str, Any], ordinary: bool = False) -> dict[str, Any]:
    walked = _walk(categories, facts)
    if not walked:
        return {"category": "未归类", "subclass": "未归类", "rw": None, "rw_text": ""}
    node = walked[-1]
    weight = ((node.get("then") or {}).get("risk_weight") or {})
    kind = weight.get("type")
    rw = weight.get("value") if kind == "const" else None
    text = weight.get("display") or ""
    if not ordinary and kind in {"counterparty", "max"}:
        other = _classify(_ordinary_categories(categories), _ordinary_facts(facts), True)
        other_rw = other.get("rw")
        if kind == "counterparty":
            rw = other_rw
            text = "按交易对手" if rw is None else _percent(rw)
        else:
            floor = weight.get("with") if isinstance(weight.get("with"), (int, float)) else 0.9
            rw = floor if other_rw is None else max(float(floor), float(other_rw))
            text = _percent(rw)
    return {
        "category": walked[0].get("label") or "",
        "subclass": node.get("label") or "",
        "path": " → ".join(str(item.get("label") or "") for item in walked if item.get("label")),
        "rw": float(rw) if isinstance(rw, (int, float)) else None,
        "rw_text": text,
    }


def _trace_text(hit: dict, facts: dict) -> str:
    parts = []
    path = str(hit.get("path") or "").strip()
    if path:
        parts.append(path)
    reason = str(facts.get("_model_reason") or "").strip()
    if reason:
        parts.append("模型：" + reason[:80])
    return "。".join(parts)[:180]


def _ordinary_categories(categories: list) -> list:
    return categories


def _ordinary_facts(facts: dict[str, Any]) -> dict[str, Any]:
    next_facts = dict(facts)
    next_facts["instrument_type"] = "claim"
    next_facts["treat_as_equity"] = "no"
    next_facts["is_defaulted"] = "no"
    next_facts["is_subordinated"] = "no"
    next_facts["has_property_collateral"] = "no"
    next_facts.pop("property_uses", None)
    next_facts.pop("cashflow_dependent", None)
    next_facts.pop("ltv", None)
    return next_facts


def _percent(number: float) -> str:
    points = number * 100
    if abs(points - round(points)) < 1e-9:
        return str(int(round(points))) + "%"
    return f"{points:.2f}".rstrip("0").rstrip(".") + "%"


def _walk(nodes: list, facts: dict[str, Any]) -> list[dict] | None:
    active = []
    skipped_would = False
    for node in nodes or []:
        if _skipped(node, facts):
            if _when(node.get("match"), facts):
                skipped_would = True
            continue
        active.append(node)
    when_hit = False
    for node in active:
        if node.get("otherwise"):
            if when_hit or skipped_would:
                continue
        elif not _when(node.get("match"), facts):
            continue
        else:
            when_hit = True
        children = node.get("children") or []
        if children:
            child = _walk(children, facts)
            if child == []:
                when_hit = False
                continue
            if child:
                return [node, *child]
            return [node]
        return [node]
    if not when_hit and skipped_would:
        return []
    return None


def _skipped(node: dict, facts: dict[str, Any]) -> bool:
    skip = node.get("skip") or {}
    if not skip.get("field"):
        return False
    return _cond({"field": skip.get("field"), "op": skip.get("op") or "eq", "value": skip.get("match")}, facts)


def _when(match: object, facts: dict[str, Any]) -> bool:
    if not isinstance(match, dict):
        return False
    if isinstance(match.get("all"), list):
        return bool(match["all"]) and all(_item(item, facts) for item in match["all"])
    if isinstance(match.get("any"), list):
        return bool(match["any"]) and any(_item(item, facts) for item in match["any"])
    if match.get("field"):
        return _cond(match, facts)
    return False


def _item(item: object, facts: dict[str, Any]) -> bool:
    if not isinstance(item, dict):
        return False
    if "all" in item or "any" in item:
        return _when(item, facts)
    return _cond(item, facts)


def _cond(cond: dict, facts: dict[str, Any]) -> bool:
    field = cond.get("field")
    if not field or field not in facts:
        return False
    actual = facts[field]
    if actual is None or actual == "":
        return False
    expected = cond.get("value") if cond.get("value") is not None else cond.get("match")
    op = cond.get("op") or "eq"
    if op == "eq":
        return actual == expected or str(actual) == str(expected)
    if op == "neq":
        return str(actual) != str(expected)
    if op == "in":
        pool = expected if isinstance(expected, list) else [expected]
        return any(actual == item or str(actual) == str(item) for item in pool)
    if op == "contains":
        pool = actual if isinstance(actual, list) else str(actual).split(",")
        return any(str(item) == str(expected) for item in pool)
    try:
        left = float(actual)
        right = float(expected)
    except (TypeError, ValueError):
        return False
    if op == "gt":
        return left > right
    if op == "gte":
        return left >= right
    if op == "lt":
        return left < right
    if op == "lte":
        return left <= right
    return False
