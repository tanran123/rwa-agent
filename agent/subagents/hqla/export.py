"""导出 9 表审计底稿（openpyxl）。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import CATEGORY_LABEL, describe_when, grade_text, issuer_category_of, level_def, round2, status_text
from .lists import LIST_DEFS

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1A5F4A")


def _sheet(book, title, rows, widths=None):
    sheet = book.create_sheet(title=title)
    for row in rows:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    if widths:
        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def export_workbook(rows, agg, recon, sample, cfg, meta, caps, lists):
    book = Workbook()
    book.remove(book.active)

    gate1_list = (cfg.get("gate1") or {}).get("checks") or []
    sec_rows = [r for r in rows if not r.get("isAdjustment")]

    # 01 输入明细
    data = [
        ["债券明细输入（单位：港元）"],
        ["数据来源", (meta or {}).get("source") or "用户导入"],
        ["基准日", cfg.get("baseDate")],
        [],
        ["行号", "申报档位", "组合与证券代码", "币种组", "申报金额(HKD)", "外部信息匹配", "关卡一", "判定档位", "命中规则", "综合结论"],
    ]
    for r in rows:
        gate1 = r.get("gate1")
        data.append([
            r.get("lineNo"), r.get("itemRaw"), r.get("secId"), r.get("ccyGroup"),
            round2(r.get("amount")) if r.get("amount") is not None else None,
            "手工调整行" if r.get("isAdjustment") else (r.get("matchBy") if r.get("ref") else "未匹配"),
            ("不适用" if not gate1 else ("不合格" if gate1["status"] == "fail" else ("通过（部分未执行）" if gate1["status"] == "pass_with_note" else "通过"))),
            level_def(cfg, r["judged"]["level"])["label"], r["judged"].get("ruleLabel"), status_text(r.get("status")),
        ])
    _sheet(book, "01-输入明细", data, [10, 14, 34, 10, 18, 18, 16, 14, 28, 12])

    # 02 关卡判定明细
    head = ["行号", "证券代码", "债券名称", "发行人类型", "债项评级", "风险权重", "30天跌幅"]
    head += [f"关卡一·{c.get('label') or c.get('code')}" for c in gate1_list]
    head += ["关卡一结论", "命中规则", "关卡二结论", "最终档位"]
    data = [["HQLA 判定：关卡一（通用合格性）× 关卡二（分档）"], ["状态：通过 / 不合格 / 未执行；数据来自内部基础参数清单与外部债券信息"], [], head]
    for r in sec_rows:
        by_code = {c["code"]: c for c in ((r.get("gate1") or {}).get("checks") or [])}
        cells = []
        for c in gate1_list:
            x = by_code.get(c["code"])
            cells.append("—" if not x else ("通过" if x["status"] == "pass" else ("不合格" if x["status"] == "fail" else "未执行")))
        ref = r.get("ref") or {}
        gate1, gate2 = r.get("gate1"), r.get("gate2") or {}
        data.append([
            r.get("lineNo"), r.get("secId"), ref.get("name") or "",
            CATEGORY_LABEL.get((r.get("ctx") or {}).get("category") or "unknown"),
            ref.get("rating") or "", "" if r.get("rw") is None else f"{r['rw']}%",
            "" if r.get("price30") is None else f"{r['price30']}%",
            *cells,
            ("—" if not gate1 else ("不合格：" + gate1["summary"] if gate1["status"] == "fail" else "通过")),
            (f"{gate2['ruleId']} {gate2['ruleLabel']}" if gate2.get("ruleId") else (gate2.get("ruleLabel") or "—")),
            ("命中" if gate2.get("hit") else "未命中") if gate2 else "—",
            level_def(cfg, r["judged"]["level"])["label"],
        ])
    _sheet(book, "02-关卡判定明细", data)

    # 03 四项核对
    data = [
        ["逐笔四项核对：分类 / 估值 / 期限 / 评级"],
        ["判定标准来自「判定标准配置」；状态：通过 / 关注 / 异常"],
        [],
        ["行号", "证券代码", "债券名称", "发行人类型", "评级", "到期日", "剩余天数", "币种组",
         "申报档位", "判定档位", "分类核对", "分类结论", "估值核对", "估值结论",
         "期限核对", "期限结论", "评级核对", "评级结论", "综合"],
    ]
    for r in sec_rows:
        ref = r.get("ref") or {}
        checks = r.get("checks") or {}
        data.append([
            r.get("lineNo"), r.get("secId"), ref.get("name") or "",
            CATEGORY_LABEL.get(issuer_category_of(ref)) if ref else "",
            ref.get("rating") or "", ref.get("mat") or "", r.get("remainingDays"), r.get("ccyGroup"),
            r.get("itemRaw"), level_def(cfg, r["judged"]["level"])["label"],
            status_text((checks.get("classification") or {}).get("status")), (checks.get("classification") or {}).get("text"),
            status_text((checks.get("valuation") or {}).get("status")), (checks.get("valuation") or {}).get("text"),
            status_text((checks.get("maturity") or {}).get("status")), (checks.get("maturity") or {}).get("text"),
            status_text((checks.get("rating") or {}).get("status")), (checks.get("rating") or {}).get("text"),
            status_text(r.get("status")),
        ])
    _sheet(book, "03-四项核对", data)

    # 04 汇总重编
    ccy_codes = agg["ccyCodes"]
    head = ["档位", "折扣系数", "本金合计(千元)"] + [f"{c} 本金(千元)" for c in ccy_codes] + [f"{c} 加权(千元)" for c in ccy_codes] + ["加权合计(千元)"]
    data = [["HQLA Compilation 口径重编（单位：千港元）"], ["加总方式：按档位 × 币种组，逐币种乘折扣系数后加总"], [], head]
    for lv in ("L1", "L2A", "L2B"):
        definition = level_def(cfg, lv)
        line = [definition["label"], definition["weight"], round2(agg["matrix"][lv]["totalAmount"] / 1000)]
        line += [round2(agg["matrix"][lv]["ccy"][c]["amount"] / 1000) for c in ccy_codes]
        line += [round2(agg["matrix"][lv]["ccy"][c]["weighted"] / 1000) for c in ccy_codes]
        line.append(round2(agg["matrix"][lv]["totalWeighted"] / 1000))
        data.append(line)
    data.append(["合计", "", round2(agg["totals"]["amount"] / 1000)]
                + [round2(agg["totals"]["ccy"][c]["amount"] / 1000) for c in ccy_codes]
                + [round2(agg["totals"]["ccy"][c]["weighted"] / 1000) for c in ccy_codes]
                + [round2(agg["totals"]["weighted"] / 1000)])
    _sheet(book, "04-汇总重编", data)

    # 05 对账
    if recon:
        data = [
            ["与客户申报的加权金额对账（单位：千港元）"],
            ["DTT 重编列由本系统按明细加总；CYB 列取自客户申报（HQLA Compilation）"],
            [],
            ["档位", "币种组", "DTT 重编", "客户申报", "差异", "结论"],
        ]
        for r in recon["rows"]:
            data.append([level_def(cfg, r["level"])["label"], r["ccyLabel"], round2(r["dtt"]), round2(r["cyb"]), round2(r["diff"]), "一致" if r["status"] == "ok" else "存在差异"])
        data += [[], ["重编加权合计", round2(recon["summary"]["dttTotal"])], ["申报加权合计", round2(recon["summary"]["cybTotal"])],
                 ["差异", round2(recon["summary"]["diff"])], ["差异率", f"{recon['summary']['diffPct'] * 100:.4f}%"]]
        _sheet(book, "05-与申报数对账", data)

    # 06 组合上限
    if caps:
        data = [
            ["关卡三：Level 2 / Level 2B 组合上限测算（单位：港元）"],
            [f"上限：Level 2A + 2B ≤ {caps['caps']['l2Pct'] * 100:.0f}% · Level 2B ≤ {caps['caps']['l2bPct'] * 100:.0f}%（以最终 HQLA 为基数）"],
            [],
            ["项目", "调整前", "调整后", "剔除金额"],
            ["Level 1", round2(caps["before"]["l1"]), round2(caps["after"]["l1"]), 0],
            ["Level 2A", round2(caps["before"]["l2a"]), round2(caps["after"]["l2a"]), round2(caps["before"]["l2a"] - caps["after"]["l2a"])],
            ["Level 2B", round2(caps["before"]["l2b"]), round2(caps["after"]["l2b"]), round2(caps["before"]["l2b"] - caps["after"]["l2b"])],
            ["HQLA 合计", round2(caps["before"]["total"]), round2(caps["after"]["total"]), round2(caps["before"]["total"] - caps["after"]["total"])],
            [],
            ["Level 2 占比", f"{caps['ratios']['l2Before'] * 100:.2f}%", f"{caps['ratios']['l2After'] * 100:.2f}%", "触及上限" if caps["binding"]["l2"] else "未触及"],
            ["Level 2B 占比", f"{caps['ratios']['l2bBefore'] * 100:.2f}%", f"{caps['ratios']['l2bAfter'] * 100:.2f}%", "触及上限" if caps["binding"]["l2b"] else "未触及"],
        ]
        _sheet(book, "06-组合上限", data)

    # 07 抽样底稿
    if sample:
        data = [
            ["抽样方法与可复现信息"],
            ["抽样方式", sample.get("method")],
            ["总体笔数", sample.get("universe")],
            ["样本笔数", sample.get("size")],
            ["覆盖率", f"{(sample.get('coverage') or 0) * 100:.2f}%"],
            ["随机种子", sample.get("seed")],
            [],
            ["分层", "总体笔数", "抽取笔数"],
        ]
        for item in sample.get("strata") or []:
            data.append(["档位未识别" if item["level"] == "UNK" else level_def(cfg, item["level"])["label"], item["population"], item["sample"]])
        data += [[], ["序号", "证券代码", "申报档位", "关卡一", "判定档位", "金额(HKD)", "综合结论"]]
        for i, r in enumerate(sample.get("rows") or []):
            gate1 = r.get("gate1")
            data.append([
                i + 1, r.get("secId"), r.get("itemRaw"),
                ("不合格" if gate1["status"] == "fail" else "通过") if gate1 else "—",
                level_def(cfg, r["judged"]["level"])["label"], r.get("amount"), status_text(r.get("status")),
            ])
        _sheet(book, "07-抽样底稿", data)

    # 08 内部清单摘要
    data = [["内部基础参数清单（关卡一判断依据）"], [], ["清单", "笔数", "匹配键", "说明"]]
    for definition in LIST_DEFS:
        count = len((lists or {}).get(definition["code"]) or [])
        data.append([definition["label"], count, " / ".join(definition.get("match") or []), definition.get("note")])
    _sheet(book, "08-内部清单摘要", data, [28, 10, 22, 60])

    # 09 配置快照
    data = [
        ["判定标准与折扣系数配置快照"],
        ["标准版本", cfg.get("standardVersion") or "v1.0"],
        ["生效日", cfg.get("effectiveFrom") or cfg.get("baseDate")],
        ["锁定状态", "已锁定" if cfg.get("locked") else "未锁定"],
        ["基准日", cfg.get("baseDate")], [],
        ["档位", "折扣系数", "最低评级门槛", "监管条款", "说明"],
    ]
    for item in cfg.get("levels") or []:
        data.append([item["label"], item["weight"], "无" if item.get("minGrade") is None else grade_text(item["minGrade"]), item.get("clause"), item.get("note")])
    data += [[], ["关卡一检查项", "启用", "说明"]]
    for item in gate1_list:
        data.append([item.get("label") or item.get("code"), "是" if item.get("enabled") else "否", ""])
    data += [[], ["规则编号", "启用", "目标档位", "规则说明", "命中条件", "监管条款"]]
    for rule in cfg.get("rules") or []:
        data.append([rule["id"], "否" if rule.get("enabled") is False else "是", level_def(cfg, rule["target"])["label"], rule["label"], describe_when(rule.get("when")), rule.get("clause")])
    gate3 = cfg.get("gate3") or {}
    data += [[], ["关卡三上限", f"Level 2 ≤ {float(gate3.get('level2CapPct') or 0.4) * 100:.0f}%", f"Level 2B ≤ {float(gate3.get('level2bCapPct') or 0.15) * 100:.0f}%"]]
    _sheet(book, "09-配置快照", data, [28, 12, 20, 50, 60, 20])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
