"""Flask 与 HQLA 引擎之间的编排层：解析上传、计算、导出。"""

from __future__ import annotations

import base64
import csv
import io
import json

from .aggregate import aggregate, apply_caps, reconcile
from .engine import cfg_default, round2, sec_key
from .export import export_workbook as build_workbook
from .imports import detect_columns, detect_ref_columns, normalize_item, parse_compilation_rows, parse_input_matrix, rows_from_ref_matrix
from .judge import analyze_rules, run
from .lists import LIST_CODES, LIST_DEFS, LIST_LABEL, build_index, list_hit, list_template, parse_list_matrix
from .sampling import plan_sample


# ---------------------------------------------------------------- 上传解析

def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return value
    return str(value)


def _read_xlsx(content, sheet_name=None):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = []
    for name in book.sheetnames:
        if sheet_name and name != sheet_name:
            continue
        sheet = book[name]
        matrix = []
        for row in sheet.iter_rows(values_only=True):
            matrix.append([_clean(cell) for cell in row])
        sheets.append({"name": name, "matrix": matrix})
    book.close()
    return sheets


def _read_delimited(text):
    sample = text[:4096]
    delimiter = ","
    for candidate in (",", "\t", ";", "|"):
        if sample.count(candidate) > sample.count(delimiter):
            delimiter = candidate
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell for cell in row] for row in reader]


def _sheets_from_payload(payload):
    content = payload.get("content")
    filename = str(payload.get("filename") or "")
    if payload.get("matrix"):
        return [{"name": payload.get("sheet") or "上传", "matrix": payload["matrix"]}]
    if not content:
        raise ValueError("没有上传文件")
    raw = base64.b64decode(content)
    if raw[:2] == b"PK" or filename.lower().endswith((".xlsx", ".xlsm")):
        return _read_xlsx(raw, payload.get("sheet") or None)
    text = raw.decode("utf-8-sig", errors="replace")
    return [{"name": filename or "上传", "matrix": _read_delimited(text)}]


def inspect(payload):
    """返回每张表的矩阵与列识别结果，供前端确认列映射。"""
    kind = payload.get("kind") or "input"
    sheets = _sheets_from_payload(payload)
    out = []
    for sheet in sheets:
        matrix = sheet["matrix"]
        if kind == "input":
            det = detect_columns(matrix)
        elif kind == "reference":
            det = detect_ref_columns(matrix)
        elif kind == "compilation":
            det = {"row": -1, "score": 0, "map": {}}
        else:
            det = parse_list_matrix(matrix)
            det = {"row": det["headerRow"], "score": len(det["map"]), "map": det["map"]}
        out.append({"name": sheet["name"], "matrix": matrix, "det": det, "rows": len(matrix)})
    best = max(range(len(out)), key=lambda i: out[i]["det"].get("score", 0)) if out else 0
    return {"sheets": out, "best": best}


def parse(payload):
    """按列映射把矩阵转成规范行。"""
    kind = payload.get("kind") or "input"
    sheets = _sheets_from_payload(payload)
    sheet = sheets[0] if len(sheets) == 1 else next((s for s in sheets if s["name"] == payload.get("sheet")), sheets[0])
    matrix = sheet["matrix"]
    det = payload.get("det") or {}
    if not det.get("map") and kind == "input":
        det = detect_columns(matrix)
    if not det.get("map") and kind == "reference":
        det = detect_ref_columns(matrix)

    if kind == "input":
        rows = parse_input_matrix(matrix, det)
    elif kind == "reference":
        rows = rows_from_ref_matrix(matrix, det)
    elif kind == "compilation":
        rows = parse_compilation_rows(matrix)
    elif kind == "list":
        rows = parse_list_matrix(matrix)["rows"]
    else:
        raise ValueError(f"未知的解析类型：{kind}")

    return {"kind": kind, "source": sheet["name"], "rows": rows}


def list_meta():
    return {
        "defs": [{"code": d["code"], "label": d["label"], "match": d.get("match") or [], "note": d.get("note") or ""} for d in LIST_DEFS],
        "templates": {code: list_template(code) for code in LIST_CODES},
        "cfg": cfg_default(),
    }


# ---------------------------------------------------------------- 计算

def _normalize_lists(raw):
    out = {}
    for code in LIST_CODES:
        value = (raw or {}).get(code)
        if isinstance(value, dict):
            out[code] = value.get("rows") or []
        elif isinstance(value, list):
            out[code] = value
        else:
            out[code] = []
    return out


def _normalize_reference(raw):
    if isinstance(raw, dict):
        return raw.get("rows") or []
    return raw or []


_ITEM_CODES = {"L1", "L2A", "L2B", "NA", "UNK"}


def normalize_input_rows(rows):
    """把原始明细行补齐 item / key / ccyGroupRaw，便于演示数据与 Koa 存储共用。"""
    out = []
    for i, raw in enumerate(rows or []):
        item_raw = raw.get("itemRaw")
        if item_raw in (None, ""):
            item_raw = raw.get("item")
        item = raw.get("item")
        if item not in _ITEM_CODES:
            item = normalize_item(item_raw)
        sec_id = raw.get("secId")
        key = raw.get("key") or sec_key(sec_id)
        out.append({
            "lineNo": raw.get("lineNo") or (i + 1),
            "itemRaw": "" if item_raw is None else str(item_raw),
            "item": item,
            "secId": "" if sec_id is None else str(sec_id).strip(),
            "key": key,
            "ccyGroupRaw": raw.get("ccyGroupRaw") or raw.get("ccyGroup") or "",
            "amount": raw.get("amount"),
        })
    return out


def normalize_cfg(raw):
    base = cfg_default()
    if isinstance(raw, dict):
        for key, value in raw.items():
            if key in ("gate1", "gate2", "gate3", "checks", "sampling", "reconciliation", "external") and isinstance(value, dict):
                merged = dict(base.get(key) or {})
                merged.update(value)
                base[key] = merged
            else:
                base[key] = value
    return base


def compute(payload):
    task = payload.get("task") or {}
    cfg = normalize_cfg(task.get("cfg"))
    lists = _normalize_lists(task.get("lists"))
    reference = _normalize_reference(payload.get("reference"))
    input_rows = normalize_input_rows((task.get("input") or {}).get("rows"))
    reported = task.get("reported") or []
    sample = task.get("sample")

    if not input_rows:
        return {
            "empty": True,
            "cfg": cfg,
            "listInfo": {code: len(lists.get(code) or []) for code in LIST_CODES},
            "stats": None,
        }

    dataset = run(input_rows, cfg, lists, reference)
    agg = aggregate(dataset["rows"], cfg, "reported")
    agg_judged = aggregate(dataset["rows"], cfg, "judged")
    caps = apply_caps(agg_judged, cfg)

    by_index = {row["i"]: row for row in dataset["rows"]}
    if sample and sample.get("indices"):
        sample_rows = [by_index[i] for i in sample["indices"] if i in by_index]
    else:
        sample_rows = dataset["rows"]

    recon = reconcile(agg, reported, cfg) if reported else None

    return {
        "empty": False,
        "rows": dataset["rows"],
        "stats": dataset["stats"],
        "byCheck": dataset["byCheck"],
        "cfgSnapshot": dataset["cfgSnapshot"],
        "runAt": dataset["runAt"],
        "agg": agg,
        "aggJ": agg_judged,
        "caps": caps,
        "recon": recon,
        "sampleRows": sample_rows,
        "listInfo": {code: len(lists.get(code) or []) for code in LIST_CODES},
        "analyzeRules": analyze_rules(cfg),
    }


def sample_plan(payload):
    task = payload.get("task") or {}
    cfg = normalize_cfg(task.get("cfg"))
    lists = _normalize_lists(task.get("lists"))
    reference = _normalize_reference(payload.get("reference"))
    input_rows = normalize_input_rows((task.get("input") or {}).get("rows"))
    dataset = run(input_rows, cfg, lists, reference)
    plan = plan_sample(dataset["rows"], cfg)
    plan["computedAt"] = dataset["runAt"]
    return plan


def export(payload):
    task = payload.get("task") or {}
    cfg = normalize_cfg(task.get("cfg"))
    lists = _normalize_lists(task.get("lists"))
    reference = _normalize_reference(payload.get("reference"))
    input_rows = normalize_input_rows((task.get("input") or {}).get("rows"))
    dataset = run(input_rows, cfg, lists, reference)
    agg = aggregate(dataset["rows"], cfg, "reported")
    agg_judged = aggregate(dataset["rows"], cfg, "judged")
    caps = apply_caps(agg_judged, cfg)
    reported = task.get("reported") or []
    recon = reconcile(agg, reported, cfg) if reported else None
    sample = task.get("sample")
    by_index = {row["i"]: row for row in dataset["rows"]}
    sample_rows = [by_index[i] for i in sample["indices"] if i in by_index] if sample and sample.get("indices") else dataset["rows"]
    sample_payload = dict(sample or {})
    sample_payload["rows"] = sample_rows
    return build_workbook(
        rows=dataset["rows"], agg=agg, recon=recon,
        sample=sample_payload if sample else None,
        cfg=cfg, meta=(task.get("input") or {}).get("meta"), caps=caps, lists=lists,
    )


def list_hits(payload):
    """某类清单命中本任务明细的笔数（前端第 2 步展示覆盖）。"""
    task = payload.get("task") or {}
    code = payload.get("code")
    lists = _normalize_lists(task.get("lists"))
    reference = _normalize_reference(payload.get("reference"))
    index = build_index(lists)
    ref_index = {str(r.get("id")): r for r in reference if r.get("id")}
    rows = normalize_input_rows((task.get("input") or {}).get("rows"))
    hits = []
    for raw in rows:
        ref = ref_index.get(raw.get("key") or "")
        if not ref and raw.get("key"):
            ref = next((r for r in reference if r.get("isin") and r.get("isin") == raw.get("key")), None)
        if list_hit(index, code, raw.get("key"), (ref or {}).get("isin"), (ref or {}).get("issuer")):
            hits.append(raw.get("lineNo"))
    return {"code": code, "label": LIST_LABEL.get(code, code), "hits": hits, "count": len(hits)}
