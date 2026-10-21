"""从信用风险底稿读出已经算好的风险加权资产。单位沿用底稿的千澳门元。"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

def _paper_path() -> Path:
    env = os.environ.get("RWA_WORKING_PAPER", "").strip()
    candidates = []
    if env:
        candidates.append(Path(env))
    candidates.extend([
        Path("/Users/tanran/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/wxid_g6n1bdg188y022_3ba9/temp/drag/Credit Risk Working 202609(1).xlsx"),
        Path.home() / "Desktop/feiliu/Credit Risk Working 202609(1).xlsx",
    ])
    for path in candidates:
        if path.is_file():
            return path
    return candidates[0]


PAPER_PATH = _paper_path()

SECTION_ZH = {
    "a": "主权",
    "b": "公共部门",
    "c": "多边开发银行",
    "d": "银行",
    "e": "公司",
    "f": "监管零售",
    "g": "住房抵押",
    "h": "现金及等价物",
    "i": "逾期贷款",
    "j": "其他表内",
}


def _num(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "—"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _label(row: tuple[Any, ...]) -> str:
    if len(row) < 2 or row[1] is None:
        return ""
    return str(row[1]).strip()


def _yi(thousand: float) -> float:
    return thousand / 100000.0


def load_working_paper(path: Path | None = None) -> dict[str, Any]:
    source = Path(path) if path else PAPER_PATH
    if not source.is_file():
        raise FileNotFoundError("找不到信用风险底稿：" + str(source))
    import openpyxl

    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        lines = []
        lines.extend(_part_one(book["Part I"]))
        lines.extend(_part_two(book["Part II"]))
        lines.extend(_part_three(book["Part III"]))
    finally:
        book.close()
    rwa_thousand = sum(item["rwa_thousand"] for item in lines)
    ead_thousand = sum(item["ead_thousand"] for item in lines)
    return {
        "file": source.name,
        "path": str(source),
        "as_of": "20260930",
        "unit": "MOP '000",
        "ead_thousand": ead_thousand,
        "rwa_thousand": rwa_thousand,
        "ead_yi": _yi(ead_thousand),
        "rwa_yi": _yi(rwa_thousand),
        "lines": lines,
    }


def _part_one(ws) -> list[dict[str, Any]]:
    section = "表内"
    detail = ""
    found = []
    for row in ws.iter_rows(max_col=12, values_only=True):
        label = _label(row)
        if not label or label in {"PART I", "TOTAL"}:
            continue
        low = label.lower()
        if low.startswith("risk weight"):
            principal = _num(row[2]) or 0.0
            ead = _num(row[5])
            if ead is None:
                ead = principal
            weight = _num(row[6])
            rwa = _num(row[7])
            if weight is None:
                continue
            if rwa is None:
                rwa = ead * weight
            if abs(ead) < 0.5 and abs(rwa) < 0.5:
                continue
            title = section if not detail else (section + " · " + detail)
            found.append(_line("Part I", title, label, ead, weight, rwa))
            continue
        if low in {"subtotal", "sum"}:
            continue
        if re.match(r"^[a-j] ", low):
            section = _section_name(label)
            detail = ""
        elif re.match(r"^\d+\s", label):
            detail = label[:28]
    return found


def _part_two(ws) -> list[dict[str, Any]]:
    found = []
    for row in ws.iter_rows(max_col=12, values_only=True):
        label = _label(row)
        if not label or label.upper() == "TOTAL" or label == "PART II":
            continue
        principal = _num(row[2])
        ccf = _num(row[5])
        after = (_num(row[6]) or 0.0) + (_num(row[7]) or 0.0)
        rwa = _num(row[10])
        if principal is None and rwa is None:
            continue
        ead = after if after else (principal or 0.0) * (ccf or 0.0)
        rwa = 0.0 if rwa is None else rwa
        if abs(ead) < 0.5 and abs(rwa) < 0.5:
            continue
        weight = (rwa / ead) if ead else 0.0
        item = _line("Part II", "表外", label, ead, weight, rwa)
        item["ccf"] = ccf
        item["principal_thousand"] = principal or 0.0
        found.append(item)
    return found


def _part_three(ws) -> list[dict[str, Any]]:
    found = []
    section = "衍生工具"
    for row in ws.iter_rows(max_col=12, values_only=True):
        label = _label(row)
        if not label or label.upper() == "TOTAL" or label == "PART III":
            continue
        if label.lower() in {"subtotal", "sum"}:
            continue
        cea = (_num(row[6]) or 0.0) + (_num(row[7]) or 0.0)
        rwa = _num(row[10])
        if cea == 0 and rwa is None:
            if _num(row[2]) is None:
                section = label
            continue
        rwa = 0.0 if rwa is None else rwa
        if abs(cea) < 0.5 and abs(rwa) < 0.5:
            continue
        weight = (rwa / cea) if cea else 0.0
        item = _line("Part III", section, label, cea, weight, rwa)
        found.append(item)
    return found


def _section_name(label: str) -> str:
    key = label[:1].lower()
    gloss = SECTION_ZH.get(key)
    if gloss and len(label) > 2 and label[1] == " ":
        return gloss
    return label[:40]


def _line(part: str, section: str, label: str, ead: float, weight: float, rwa: float) -> dict[str, Any]:
    weight_text = _weight_text(weight)
    return {
        "part": part,
        "section": section,
        "label": label,
        "weight": weight,
        "weight_text": weight_text,
        "ead_thousand": ead,
        "rwa_thousand": rwa,
        "ead_yi": _yi(ead),
        "rwa_yi": _yi(rwa),
    }


def _weight_text(weight: float) -> str:
    pct = weight * 100
    if abs(pct - round(pct)) < 0.05:
        return str(int(round(pct))) + "%"
    return f"{pct:.2f}%"
